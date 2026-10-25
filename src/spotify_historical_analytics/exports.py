from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import DatabaseManager


METRICS = {
    "songs": (
        "Total chart streams: sum of imported US chart observations for each track "
        "in the selected chart type and date range."
    ),
    "artists": (
        "Charting-song streams: streams credited in full to each supplied artist ID; "
        "collaborating artist totals can overlap."
    ),
    "albums": (
        "Album charting-track streams: imported chart streams for tracks mapped to the "
        "supplied album ID; these are not comprehensive album totals."
    ),
}

RANKING_COLUMNS = {
    "songs": [
        "Rank", "Song", "Artist", "Total chart streams", "Best position",
        "Chart appearances", "Track identity",
    ],
    "artists": [
        "Rank", "Artist", "Artist ID", "Charting-song streams", "Attribution",
        "Charting tracks", "Chart appearances", "Best position",
        "First appearance", "Latest appearance",
    ],
    "albums": [
        "Rank", "Album", "Album artist", "Album ID", "Album charting-track streams",
        "Mapped charting tracks", "Chart appearances", "Best position",
        "First appearance", "Latest appearance",
    ],
}


def _ranking_rows(
    db: DatabaseManager,
    category: str,
    start_date: str | None,
    end_date: str | None,
    chart_type: str,
    limit: int,
    search: str | None,
) -> list[dict[str, Any]]:
    if category == "songs":
        rows = db.get_top_tracks(
            limit, start_date, end_date, chart_type=chart_type, search=search
        )
        return [
            {
                "Rank": rank,
                "Song": row["track_name"],
                "Artist": row["artist_names"],
                "Total chart streams": int(row["total_streams"]),
                "Best position": row["best_position"],
                "Chart appearances": int(row["chart_appearances"]),
                "Track identity": row["track_key"],
            }
            for rank, row in enumerate(rows, 1)
        ]
    if category == "artists":
        rows = db.get_artist_summary(
            start_date, end_date, chart_type=chart_type, limit=limit, search=search
        )
        return [
            {
                "Rank": rank,
                "Artist": row["artist_name"],
                "Artist ID": row["artist_id"] or "",
                "Charting-song streams": int(row["charting_song_streams"]),
                "Attribution": row["attribution"],
                "Charting tracks": int(row["track_count"]),
                "Chart appearances": int(row["chart_appearances"]),
                "Best position": row["best_position"],
                "First appearance": row["first_appearance"],
                "Latest appearance": row["latest_appearance"],
            }
            for rank, row in enumerate(rows, 1)
        ]
    if category == "albums":
        rows = db.get_album_rankings(
            start_date, end_date, chart_type, limit=limit, search=search
        )
        return [
            {
                "Rank": rank,
                "Album": row["album_name"],
                "Album artist": row["album_artist"],
                "Album ID": row["album_id"],
                "Album charting-track streams": int(row["charting_track_streams"]),
                "Mapped charting tracks": int(row["mapped_track_count"]),
                "Chart appearances": int(row["chart_appearances"]),
                "Best position": row["best_position"],
                "First appearance": row["first_appearance"],
                "Latest appearance": row["latest_appearance"],
            }
            for rank, row in enumerate(rows, 1)
        ]
    raise ValueError(f"Unsupported export category: {category}")


def _report_metadata(
    db: DatabaseManager,
    category: str,
    rows: list[dict[str, Any]],
    start_date: str | None,
    end_date: str | None,
    chart_type: str,
    search: str | None,
    limit: int,
) -> tuple[list[tuple[str, str]], list[dict[str, Any]]]:
    bounds = db.get_available_date_bounds(chart_type=chart_type)
    effective_start = start_date or bounds["earliest"]
    effective_end = end_date or bounds["latest"]
    summary = db.get_summary(start_date, end_date, chart_type=chart_type)
    inventory = db.get_data_inventory()
    selected_sources = [
        source
        for source in inventory["sources"]
        if source["country"] == "US"
        and source["chart_type"] == chart_type
        and not source["is_demo"]
        and (effective_start is None or source["latest"] >= effective_start)
        and (effective_end is None or source["earliest"] <= effective_end)
    ]
    source_names = sorted({str(source["source"]) for source in selected_sources})
    observed_coverage = (
        f"{summary['coverage_percentage']:.2f}% (observed only; not proof of completeness)"
    )
    metadata = [
        ("Category", category.title()),
        ("Market", "United States (US)"),
        ("Chart type", chart_type),
        ("Date interval", f"{effective_start or 'No observations'} to {effective_end or 'No observations'}"),
        ("Date filter", f"{start_date or 'All available'} to {end_date or 'All available'}"),
        ("Search filter", search or "None"),
        ("Ranking selection", f"Top {limit} matching results"),
        ("Metric definition", METRICS[category]),
        ("Observed coverage", observed_coverage),
        ("Observed chart dates", str(summary["chart_dates_with_data"])),
        ("Rows exported", str(len(rows))),
        ("Source names", "; ".join(source_names) or "No matching real-data sources"),
        ("Source completeness", "Unverified; missing dates are not treated as zero streams."),
        ("Synthetic data", "Excluded from this real-data export."),
    ]
    return metadata, selected_sources


def _coverage_rows(sources: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "Source": source["source"],
            "Chart type": source["chart_type"],
            "Market": source["country"],
            "Rows in database": int(source["row_count"]),
            "Earliest observed date": source["earliest"],
            "Latest observed date": source["latest"],
            "Source attribution / license note": source["source_license"] or "Not recorded",
            "Completeness": "Unverified",
        }
        for source in sources
    ]


def _safe_spreadsheet_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    if value.lstrip(" \t\r\n")[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def export_rankings(
    db: DatabaseManager,
    category: str,
    destination: str | Path,
    start_date: str | None,
    end_date: str | None,
    chart_type: str,
    limit: int,
    search: str | None = None,
) -> int:
    db.validate_date_range(start_date, end_date)
    category = category.lower()
    path = Path(destination)
    suffix = path.suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ValueError("Exports must use a .csv or .xlsx destination.")
    rows = _ranking_rows(db, category, start_date, end_date, chart_type, limit, search)
    metadata, sources = _report_metadata(
        db, category, rows, start_date, end_date, chart_type, search, limit
    )
    if suffix == ".csv":
        _write_csv(path, RANKING_COLUMNS[category], rows, metadata, sources)
    else:
        _write_xlsx(path, RANKING_COLUMNS[category], rows, metadata, sources)
    return len(rows)


def _write_csv(
    path: Path,
    columns: list[str],
    rows: list[dict[str, Any]],
    metadata: list[tuple[str, str]],
    sources: list[dict[str, Any]],
) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(["Report metadata"])
        writer.writerows(
            [
                _safe_spreadsheet_value(key),
                _safe_spreadsheet_value(value),
            ]
            for key, value in metadata
        )
        writer.writerow([])
        writer.writerow(["Source coverage"])
        coverage_columns = list(_coverage_rows(sources)[0]) if sources else [
            "Source", "Chart type", "Market", "Rows in database", "Earliest observed date",
            "Latest observed date", "Source attribution / license note", "Completeness",
        ]
        writer.writerow(coverage_columns)
        for source in _coverage_rows(sources):
            writer.writerow(
                [_safe_spreadsheet_value(source[column]) for column in coverage_columns]
            )
        writer.writerow([])
        writer.writerow(["Rankings"])
        writer.writerow(columns)
        for row in rows:
            writer.writerow(
                [_safe_spreadsheet_value(row.get(column, "")) for column in columns]
            )


def _style_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="8B5CF6")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for column_cells in sheet.columns:
        values = [str(cell.value or "") for cell in column_cells]
        width = min(max(max(map(len, values), default=10) + 2, 12), 60)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    sheet.sheet_view.showGridLines = False


def _write_xlsx(
    path: Path,
    columns: list[str],
    rows: list[dict[str, Any]],
    metadata: list[tuple[str, str]],
    sources: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    rankings = workbook.active
    rankings.title = "Rankings"
    rankings.append(columns)
    for row in rows:
        rankings.append(
            [_safe_spreadsheet_value(row.get(column, "")) for column in columns]
        )
    _style_sheet(rankings)

    definitions = workbook.create_sheet("Definitions")
    definitions.append(["Field", "Value"])
    for row in metadata:
        definitions.append(
            [_safe_spreadsheet_value(value) for value in row]
        )
    _style_sheet(definitions)

    coverage = workbook.create_sheet("Coverage & Sources")
    coverage_rows = _coverage_rows(sources)
    source_columns = list(coverage_rows[0]) if coverage_rows else [
        "Source", "Chart type", "Market", "Rows in database", "Earliest observed date",
        "Latest observed date", "Source attribution / license note", "Completeness",
    ]
    coverage.append(source_columns)
    for row in coverage_rows:
        coverage.append(
            [_safe_spreadsheet_value(row[column]) for column in source_columns]
        )
    _style_sheet(coverage)
    workbook.save(path)
