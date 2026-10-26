from __future__ import annotations

import csv
import gzip
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import uuid
from typing import Callable, Iterator

from openpyxl import load_workbook

from .db import DatabaseManager


IMPORT_SESSION_SCHEMA = 2
IMPORT_ADAPTER_VERSION = 3
CHUNK_SIZE = 2_000
MAX_REJECTION_SAMPLES = 100
FIELD_ALIASES = {
    "date": ("date", "chartdate", "datevalue"),
    "track_name": ("track", "trackname", "title", "song"),
    "artist_names": ("artist", "artistnames", "artist_name", "artists"),
    "streams": ("streams", "stream", "totalstreams", "total_streams"),
    "country": ("region", "country", "market"),
    "position": ("position", "rank"),
    "url": ("url", "uri", "spotifyurl", "spotify_url"),
    "chart_type": ("chart", "chartname", "charttype", "cadence"),
    "weekly_start": ("weeklystart", "weekstart", "week_start"),
    "weekly_end": ("weeklyend", "weekend", "week_end"),
    "track_id": ("trackid", "spotifytrackid", "spotify_id"),
    "artist_id": ("artistid", "spotifyartistid"),
    "artist_name": ("mappedartistname", "resolvedartistname"),
    "album_id": ("albumid", "spotifyalbumid"),
    "album_name": ("albumname", "album"),
    "album_artist": ("albumartist", "album_artist"),
    "release_date": ("releasedate", "release_date"),
    "source_reference": ("sourcereference", "source_reference"),
}


def _normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).casefold())


def _as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _as_text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    return None


def _as_int(value: object) -> int | None:
    text = _as_text(value)
    if not text:
        return None
    if "," in text and not re.fullmatch(r"-?\d{1,3}(?:,\d{3})+(?:\.0+)?", text):
        return None
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation:
        return None
    if not number.is_finite() or number != number.to_integral_value():
        return None
    if number < -(2 ** 63) or number >= 2 ** 63:
        return None
    return int(number)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _fingerprint(path: Path) -> dict[str, int | str]:
    details = path.stat()
    return {
        "sha256": sha256_file(path),
        "size_bytes": details.st_size,
        "modified_ns": details.st_mtime_ns,
    }


@dataclass(frozen=True)
class ImportSessionSummary:
    session_id: str
    source_name: str
    file_path: str
    status: str
    current_stage: int
    processed_rows: int
    accepted_rows: int
    rejected_rows: int
    exact_duplicates: int
    conflicts: int
    unique_track_ids: int
    fallback_id_rows: int
    country_counts: dict[str, int]
    chart_type_counts: dict[str, int]
    rejection_counts: dict[str, int]
    first_date: str | None
    last_date: str | None
    coverage_by_year: dict[str, int]
    artist_mapping_rows: int
    album_mapping_rows: int
    missing_date_candidates: list[str]
    rights_status: str
    error: str | None = None


class ImportSessionStore:
    """SQLite staging store separate from production chart records."""

    def __init__(self, database_path: str | Path) -> None:
        self.database_path = Path(database_path)
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(str(self.database_path), timeout=30)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute("PRAGMA journal_mode=WAL")
        self.connection.execute("PRAGMA foreign_keys=ON")
        self._initialize()

    def close(self) -> None:
        self.connection.close()

    def _initialize(self) -> None:
        version = int(self.connection.execute("PRAGMA user_version").fetchone()[0])
        if version > IMPORT_SESSION_SCHEMA:
            raise RuntimeError(
                f"Import-session schema {version} is newer than supported "
                f"version {IMPORT_SESSION_SCHEMA}."
            )
        self.connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS import_sessions (
                session_id TEXT PRIMARY KEY,
                source_name TEXT NOT NULL,
                source_license TEXT NOT NULL,
                chart_type TEXT NOT NULL,
                file_path TEXT NOT NULL,
                file_size INTEGER NOT NULL,
                file_modified_ns INTEGER NOT NULL,
                file_sha256 TEXT,
                adapter_version INTEGER NOT NULL,
                schema_json TEXT NOT NULL DEFAULT '{}',
                headers_json TEXT NOT NULL DEFAULT '[]',
                status TEXT NOT NULL DEFAULT 'new',
                current_stage INTEGER NOT NULL DEFAULT 1,
                last_source_row INTEGER NOT NULL DEFAULT 0,
                rows_processed INTEGER NOT NULL DEFAULT 0,
                approval_json TEXT,
                rights_status TEXT NOT NULL DEFAULT 'UNVERIFIED',
                error TEXT,
                created_at TEXT NOT NULL DEFAULT (datetime('now')),
                updated_at TEXT NOT NULL DEFAULT (datetime('now'))
            );
            CREATE TABLE IF NOT EXISTS staged_chart_rows (
                session_id TEXT NOT NULL,
                source_row INTEGER NOT NULL,
                chart_date TEXT NOT NULL,
                chart_type TEXT NOT NULL,
                country TEXT NOT NULL,
                position INTEGER,
                track_id TEXT NOT NULL,
                track_name TEXT NOT NULL,
                artist_names TEXT NOT NULL,
                artist_id TEXT,
                artist_name TEXT,
                streams INTEGER NOT NULL,
                original_url TEXT,
                weekly_start TEXT,
                weekly_end TEXT,
                album_id TEXT,
                album_name TEXT,
                album_artist TEXT,
                release_date TEXT,
                source_reference TEXT,
                record_hash TEXT NOT NULL,
                PRIMARY KEY (session_id, source_row),
                FOREIGN KEY (session_id) REFERENCES import_sessions(session_id)
                    ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_staged_rows_identity
                ON staged_chart_rows(session_id, chart_date, chart_type, track_id);
            CREATE INDEX IF NOT EXISTS idx_staged_rows_hash
                ON staged_chart_rows(session_id, record_hash);
            CREATE TABLE IF NOT EXISTS import_rejection_counts (
                session_id TEXT NOT NULL,
                reason TEXT NOT NULL,
                row_count INTEGER NOT NULL,
                PRIMARY KEY (session_id, reason),
                FOREIGN KEY (session_id) REFERENCES import_sessions(session_id)
                    ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS import_rejection_samples (
                session_id TEXT NOT NULL,
                source_row INTEGER NOT NULL,
                reason TEXT NOT NULL,
                PRIMARY KEY (session_id, source_row),
                FOREIGN KEY (session_id) REFERENCES import_sessions(session_id)
                    ON DELETE CASCADE
            );
            """
        )
        session_columns = {
            row["name"]
            for row in self.connection.execute("PRAGMA table_info(import_sessions)")
        }
        if "headers_json" not in session_columns:
            self.connection.execute(
                "ALTER TABLE import_sessions ADD COLUMN headers_json TEXT NOT NULL DEFAULT '[]'"
            )
        staged_columns = {
            row["name"]
            for row in self.connection.execute("PRAGMA table_info(staged_chart_rows)")
        }
        if "artist_name" not in staged_columns:
            self.connection.execute(
                "ALTER TABLE staged_chart_rows ADD COLUMN artist_name TEXT"
            )
        self.connection.execute(f"PRAGMA user_version={IMPORT_SESSION_SCHEMA}")
        self.connection.commit()

    def create_session(
        self,
        file_path: str | Path,
        source_name: str,
        source_license: str,
        chart_type: str,
    ) -> str:
        path = Path(file_path).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Selected dataset file does not exist: {path}")
        if not path.name.casefold().endswith((".csv", ".csv.gz", ".xlsx")):
            raise ValueError("Select a CSV, CSV.GZ, or XLSX dataset.")
        if not source_name.strip():
            raise ValueError("Enter a source name before preparing the dataset.")
        normalized_chart = DatabaseManager.normalize_chart_type(chart_type)
        if normalized_chart not in {"daily", "weekly"}:
            raise ValueError("Choose an explicit daily or weekly chart cadence.")
        stat = path.stat()
        session_id = uuid.uuid4().hex
        self.connection.execute(
            """
            INSERT INTO import_sessions(
                session_id, source_name, source_license, chart_type, file_path,
                file_size, file_modified_ns, adapter_version
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                session_id,
                source_name.strip(),
                source_license.strip(),
                normalized_chart,
                str(path),
                stat.st_size,
                stat.st_mtime_ns,
                IMPORT_ADAPTER_VERSION,
            ),
        )
        self.connection.commit()
        return session_id

    def get_session(self, session_id: str) -> sqlite3.Row:
        row = self.connection.execute(
            "SELECT * FROM import_sessions WHERE session_id=?", (session_id,)
        ).fetchone()
        if row is None:
            raise KeyError(f"Unknown import session: {session_id}")
        return row

    def resumable_sessions(self) -> list[sqlite3.Row]:
        return self.connection.execute(
            """
            SELECT * FROM import_sessions
            WHERE status IN ('new', 'detected', 'paused', 'staged', 'blocked')
            ORDER BY updated_at DESC
            """
        ).fetchall()

    def discard_session(self, session_id: str) -> None:
        self.connection.execute(
            "DELETE FROM import_sessions WHERE session_id=?", (session_id,)
        )
        self.connection.commit()

    def mark_failed(self, session_id: str, error: str) -> None:
        self.connection.execute(
            """
            UPDATE import_sessions SET status='failed', error=?,
                updated_at=datetime('now') WHERE session_id=?
            """,
            (error[:2_000], session_id),
        )
        self.connection.commit()

    @staticmethod
    def _read_rows(path: Path) -> Iterator[tuple[list[str], Iterator[dict[str, object]]]]:
        if path.name.casefold().endswith((".csv", ".csv.gz")):
            opener = gzip.open if path.name.casefold().endswith(".gz") else open
            source = opener(path, "rt", encoding="utf-8-sig", newline="")
            sample = source.read(16384)
            source.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
            except csv.Error:
                dialect = csv.excel
            reader = csv.DictReader(source, dialect=dialect)
            if reader.fieldnames is None:
                source.close()
                raise ValueError("The CSV file has no header row.")

            def rows() -> Iterator[dict[str, object]]:
                try:
                    yield from reader
                finally:
                    source.close()

            yield [str(name) for name in reader.fieldnames], rows()
            return

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)
            header_values = next(values, None)
            if header_values is None:
                raise ValueError("The XLSX worksheet is empty.")
            headers = [str(value).strip() if value is not None else "" for value in header_values]

            def rows() -> Iterator[dict[str, object]]:
                try:
                    for row in values:
                        yield {
                            header: value
                            for header, value in zip(headers, row)
                            if header
                        }
                finally:
                    workbook.close()

            yield headers, rows()
        except Exception:
            workbook.close()
            raise

    def _resolve_schema(self, headers: list[str]) -> dict[str, str]:
        by_normalized = {_normalize_header(header): header for header in headers if header}
        resolved = {
            field: by_normalized[_normalize_header(alias)]
            for field, aliases in FIELD_ALIASES.items()
            for alias in aliases
            if _normalize_header(alias) in by_normalized
        }
        if self.detect_adapter(headers) == "MGD+":
            resolved.update(track_id="ID", country="market")
        return resolved

    @staticmethod
    def detect_adapter(headers: list[str]) -> str:
        from .source_adapters import MGDPLUS_CHART_COLUMNS, DHRUVIL_COLUMNS
        if set(MGDPLUS_CHART_COLUMNS).issubset(headers):
            return "MGD+"
        if set(DHRUVIL_COLUMNS).issubset(headers):
            return "Dhruvil Spotify Charts"
        return "Compatible table"

    @staticmethod
    def _validate_schema(schema: dict[str, str], headers: list[str]) -> None:
        required = ("date", "track_name", "artist_names", "streams", "country")
        missing = [name for name in required if not schema.get(name)]
        if missing:
            raise ValueError(
                "Required dataset columns were not detected: "
                + ", ".join(missing)
                + ". Required fields are date, track, artist, streams, and region/country."
            )
        invalid = [name for name, column in schema.items() if column not in headers]
        if invalid:
            raise ValueError(
                "Mapped source columns are missing: " + ", ".join(invalid)
            )
        mapped_columns = [schema[name] for name in required]
        if len(set(mapped_columns)) != len(mapped_columns):
            raise ValueError("Each required field must use a distinct source column.")

    def inspect_file(self, session_id: str) -> dict[str, object]:
        session = self.get_session(session_id)
        path = Path(session["file_path"])
        if not path.is_file():
            raise FileNotFoundError(f"The selected source file is no longer available: {path}")
        if session["adapter_version"] != IMPORT_ADAPTER_VERSION:
            raise ValueError(
                "The import adapter version changed. Start a new session to revalidate."
            )
        fingerprint = _fingerprint(path)
        if (
            fingerprint["size_bytes"] != session["file_size"]
            or fingerprint["modified_ns"] != session["file_modified_ns"]
            or (
                session["file_sha256"] is not None
                and fingerprint["sha256"] != session["file_sha256"]
            )
        ):
            self.connection.execute(
                """
                UPDATE import_sessions SET status='stale', file_sha256=?,
                    error='Source file changed; revalidation required.',
                    updated_at=datetime('now') WHERE session_id=?
                """,
                (fingerprint["sha256"], session_id),
            )
            self.connection.commit()
            raise ValueError(
                "The source file changed since this import session was prepared. "
                "Start a new session; staged validation cannot be resumed."
            )
        iterator = self._read_rows(path)
        headers, rows = next(iterator)
        close = getattr(rows, "close", None)
        if close:
            close()
        if len([header for header in headers if header]) != len(
            {header for header in headers if header}
        ):
            raise ValueError("The dataset contains duplicate column headers.")
        schema = self._resolve_schema(headers)
        with self.connection:
            self.connection.execute(
                """
                UPDATE import_sessions SET file_sha256=?, schema_json=?,
                    headers_json=?, status='detected', current_stage=4, error=NULL,
                    updated_at=datetime('now') WHERE session_id=?
                """,
                (
                    fingerprint["sha256"],
                    json.dumps(schema, sort_keys=True),
                    json.dumps(headers, ensure_ascii=False),
                    session_id,
                ),
            )
        return {
            "format": "csv.gz" if path.name.casefold().endswith(".csv.gz") else path.suffix.casefold().lstrip("."),
            "adapter": self.detect_adapter(headers),
            "headers": headers,
            "schema": schema,
            "fingerprint": fingerprint,
        }

    def session_headers(self, session_id: str) -> list[str]:
        row = self.get_session(session_id)
        return json.loads(row["headers_json"])

    def save_schema_mapping(
        self, session_id: str, schema: dict[str, str]
    ) -> None:
        session = self.get_session(session_id)
        if session["status"] not in {"detected", "paused", "staged"}:
            raise ValueError("Inspect the source file before mapping its columns.")
        path = Path(session["file_path"])
        iterator = self._read_rows(path)
        headers, rows = next(iterator)
        close = getattr(rows, "close", None)
        if close:
            close()
        self._validate_schema(schema, headers)
        with self.connection:
            self.connection.execute(
                "UPDATE import_sessions SET schema_json=?, updated_at=datetime('now') WHERE session_id=?",
                (json.dumps(schema, sort_keys=True), session_id),
            )

    @staticmethod
    def _record_hash(record: dict[str, object]) -> str:
        observation_fields = (
            "chart_date",
            "chart_type",
            "country",
            "position",
            "track_id",
            "track_name",
            "artist_names",
            "streams",
            "original_url",
            "weekly_start",
            "weekly_end",
        )
        observation = {key: record[key] for key in observation_fields}
        raw = json.dumps(
            observation, sort_keys=True, ensure_ascii=False, separators=(",", ":")
        )
        return hashlib.sha256(raw.encode("utf-8")).hexdigest()

    def prepare(
        self,
        session_id: str,
        cancel_event: object | None = None,
        progress_callback: Callable[[int, str], None] | None = None,
        schema_mapping: dict[str, str] | None = None,
    ) -> ImportSessionSummary:
        session = self.get_session(session_id)
        if session["file_sha256"] is None:
            self.inspect_file(session_id)
            session = self.get_session(session_id)
        path = Path(session["file_path"])
        if not path.is_file():
            raise FileNotFoundError(f"The selected source file is no longer available: {path}")
        if session["adapter_version"] != IMPORT_ADAPTER_VERSION:
            raise ValueError(
                "The import adapter version changed. Start a new session to revalidate."
            )
        self.connection.execute(
            "UPDATE import_sessions SET status='scanning', error=NULL WHERE session_id=?",
            (session_id,),
        )
        self.connection.commit()
        fingerprint = _fingerprint(path)
        current = self.get_session(session_id)
        if (
            fingerprint["size_bytes"] != current["file_size"]
            or fingerprint["modified_ns"] != current["file_modified_ns"]
            or (
                current["file_sha256"] is not None
                and fingerprint["sha256"] != current["file_sha256"]
            )
        ):
            self.connection.execute(
                """
                UPDATE import_sessions SET status='stale',
                    file_sha256=?, error='Source file changed; revalidation required.',
                    updated_at=datetime('now') WHERE session_id=?
                """,
                (fingerprint["sha256"], session_id),
            )
            self.connection.commit()
            raise ValueError(
                "The source file changed since this import session was prepared. "
                "Start a new session; staged validation cannot be resumed."
            )
        with self.connection:
            self.connection.execute(
                "UPDATE import_sessions SET file_sha256=? WHERE session_id=?",
                (fingerprint["sha256"], session_id),
            )

        iterator = self._read_rows(path)
        headers, rows = next(iterator)
        schema = (
            schema_mapping
            or json.loads(self.get_session(session_id)["schema_json"])
            or self._resolve_schema(headers)
        )
        self._validate_schema(schema, headers)
        self.connection.execute(
            "UPDATE import_sessions SET schema_json=? WHERE session_id=?",
            (json.dumps(schema, sort_keys=True), session_id),
        )
        self.connection.commit()
        last_source_row = int(self.get_session(session_id)["last_source_row"])
        pending_records: list[tuple[object, ...]] = []
        pending_rejections: list[tuple[str, int, str]] = []
        processed_since_commit = 0
        for source_row, raw in enumerate(rows, start=2):
            if source_row <= last_source_row:
                continue
            if cancel_event is not None and cancel_event.is_set():
                self._flush_chunk(session_id, pending_records, pending_rejections, source_row - 1)
                self.connection.execute(
                    "UPDATE import_sessions SET status='paused', updated_at=datetime('now') WHERE session_id=?",
                    (session_id,),
                )
                self.connection.commit()
                return self.summary(session_id)
            record, reason = self._normalize_row(
                raw,
                schema,
                str(session["chart_type"]),
            )
            if record is None:
                pending_rejections.append((session_id, source_row, reason))
            else:
                record_hash = self._record_hash(record)
                pending_records.append(
                    (
                        session_id,
                        source_row,
                        record["chart_date"],
                        record["chart_type"],
                        record["country"],
                        record["position"],
                        record["track_id"],
                        record["track_name"],
                        record["artist_names"],
                        record["artist_id"],
                        record["artist_name"],
                        record["streams"],
                        record["original_url"],
                        record["weekly_start"],
                        record["weekly_end"],
                        record["album_id"],
                        record["album_name"],
                        record["album_artist"],
                        record["release_date"],
                        record["source_reference"],
                        record_hash,
                    )
                )
            processed_since_commit += 1
            if processed_since_commit >= CHUNK_SIZE:
                self._flush_chunk(
                    session_id, pending_records, pending_rejections, source_row
                )
                pending_records.clear()
                pending_rejections.clear()
                processed_since_commit = 0
                last_source_row = source_row
                if progress_callback:
                    progress_callback(source_row - 1, "Validating and staging CSV/XLSX rows")
        if processed_since_commit:
            self._flush_chunk(
                session_id, pending_records, pending_rejections, source_row
                if "source_row" in locals()
                else last_source_row
            )
        self.connection.execute(
            """
            UPDATE import_sessions SET status='staged', error=NULL,
                current_stage=4, updated_at=datetime('now')
            WHERE session_id=?
            """,
            (session_id,),
        )
        self.connection.commit()
        return self.summary(session_id)

    @staticmethod
    def _normalize_row(
        raw: dict[str, object], schema: dict[str, str], selected_chart: str
    ) -> tuple[dict[str, object] | None, str]:
        def field(name: str) -> object:
            column = schema.get(name)
            return raw.get(column) if column else None

        chart_date = _parse_date(field("date"))
        if not chart_date:
            return None, "invalid_date"
        country_value = _as_text(field("country"))
        country = DatabaseManager.normalize_country(country_value)
        if not country:
            return None, "missing_country"
        if country != "US":
            return None, f"non_us_country:{country}"
        track_name = _as_text(field("track_name"))
        artist_names = _as_text(field("artist_names"))
        if not track_name or not artist_names:
            return None, "missing_track_or_artist"
        streams = _as_int(field("streams"))
        if streams is None or streams < 0:
            return None, "invalid_streams"
        adapter = ImportSessionStore.detect_adapter(list(raw))
        chart_value = _as_text(field("chart_type"))
        if adapter != "Compatible table":
            if not chart_value:
                return None, "missing_chart_category"
            if "viral" in chart_value.casefold():
                return None, "streamless_viral_chart"
            chart_value = "daily"
        chart_type = (
            DatabaseManager.normalize_chart_type(chart_value)
            if chart_value
            else selected_chart
        )
        if chart_type not in {"daily", "weekly"}:
            return None, "unsupported_chart_type"
        if chart_type != selected_chart:
            return None, "chart_type_mismatch"
        raw_track_id = _as_text(field("track_id"))
        url = _as_text(field("url"))
        if not raw_track_id and url:
            from .source_adapters import _verified_spotify_track_id
            raw_track_id = _verified_spotify_track_id(url) or ""
        if adapter != "Compatible table" and not raw_track_id:
            return None, "missing_stable_track_id"
        track_id = raw_track_id or f"{track_name}|{artist_names}"
        position = _as_int(field("position"))
        if (adapter != "Compatible table" or _as_text(field("position"))) and (position is None or position < 1):
            return None, "invalid_chart_position"
        weekly_start = _parse_date(field("weekly_start"))
        weekly_end = _parse_date(field("weekly_end"))
        if selected_chart == "weekly" and (not weekly_start or not weekly_end):
            return None, "missing_week_boundaries"
        if weekly_start and weekly_end and weekly_start > weekly_end:
            return None, "invalid_week_boundaries"
        release_date_value = field("release_date")
        release_date = _parse_date(release_date_value) if release_date_value else None
        record = {
            "chart_date": chart_date,
            "chart_type": selected_chart,
            "country": "US",
            "position": position,
            "track_id": track_id,
            "track_name": track_name,
            "artist_names": artist_names,
            "artist_id": _as_text(field("artist_id")) or None,
            "artist_name": _as_text(field("artist_name")) or None,
            "streams": streams,
            "original_url": url or None,
            "weekly_start": weekly_start,
            "weekly_end": weekly_end,
            "album_id": _as_text(field("album_id")) or None,
            "album_name": _as_text(field("album_name")) or None,
            "album_artist": _as_text(field("album_artist")) or None,
            "release_date": release_date,
            "source_reference": _as_text(field("source_reference")) or None,
        }
        if record["artist_id"] and not record["artist_name"]:
            record["artist_id"] = None
        if record["album_id"] and not record["album_name"]:
            return None, "album_name_missing_for_album_id"
        return record, ""

    def _flush_chunk(
        self,
        session_id: str,
        records: list[tuple[object, ...]],
        rejections: list[tuple[str, int, str]],
        last_source_row: int,
    ) -> None:
        with self.connection:
            self.connection.executemany(
                """
                INSERT OR IGNORE INTO staged_chart_rows(
                    session_id, source_row, chart_date, chart_type, country,
                    position, track_id, track_name, artist_names, artist_id,
                    artist_name, streams, original_url, weekly_start, weekly_end, album_id,
                    album_name, album_artist, release_date, source_reference,
                    record_hash
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                records,
            )
            sample_count = self.connection.execute(
                "SELECT COUNT(*) FROM import_rejection_samples WHERE session_id=?",
                (session_id,),
            ).fetchone()[0]
            for rejected_session, source_row, reason in rejections:
                self.connection.execute(
                    """
                    INSERT INTO import_rejection_counts(session_id, reason, row_count)
                    VALUES (?, ?, 1)
                    ON CONFLICT(session_id, reason)
                    DO UPDATE SET row_count=row_count+1
                    """,
                    (rejected_session, reason),
                )
                if sample_count < MAX_REJECTION_SAMPLES:
                    self.connection.execute(
                        """
                        INSERT OR IGNORE INTO import_rejection_samples(
                            session_id, source_row, reason
                        ) VALUES (?, ?, ?)
                        """,
                        (session_id, source_row, reason),
                    )
                    sample_count += 1
            self.connection.execute(
                """
                UPDATE import_sessions SET last_source_row=?, rows_processed=?,
                    updated_at=datetime('now') WHERE session_id=?
                """,
                (last_source_row, last_source_row - 1, session_id),
            )

    def summary(self, session_id: str) -> ImportSessionSummary:
        session = self.get_session(session_id)
        accepted = int(
            self.connection.execute(
                "SELECT COUNT(*) FROM staged_chart_rows WHERE session_id=?",
                (session_id,),
            ).fetchone()[0]
        )
        rejection_rows = self.connection.execute(
            """
            SELECT reason, row_count FROM import_rejection_counts
            WHERE session_id=? ORDER BY reason
            """,
            (session_id,),
        ).fetchall()
        rejection_counts = {row["reason"]: int(row["row_count"]) for row in rejection_rows}
        country_counts = {
            row["country"]: int(row["n"])
            for row in self.connection.execute(
                """
                SELECT country, COUNT(*) AS n FROM staged_chart_rows
                WHERE session_id=? GROUP BY country
                """,
                (session_id,),
            ).fetchall()
        }
        chart_counts = {
            row["chart_type"]: int(row["n"])
            for row in self.connection.execute(
                """
                SELECT chart_type, COUNT(*) AS n FROM staged_chart_rows
                WHERE session_id=? GROUP BY chart_type
                """,
                (session_id,),
            ).fetchall()
        }
        exact_duplicates = max(
            0,
            accepted
            - int(
                self.connection.execute(
                    """
                    SELECT COUNT(DISTINCT record_hash) FROM staged_chart_rows
                    WHERE session_id=?
                    """,
                    (session_id,),
                ).fetchone()[0]
            ),
        )
        conflict_row = self.connection.execute(
            """
            SELECT COUNT(*) AS conflicts FROM (
                SELECT chart_date, chart_type, track_id
                FROM staged_chart_rows WHERE session_id=?
                GROUP BY chart_date, chart_type, track_id
                HAVING COUNT(DISTINCT COALESCE(position, -1)) > 1
                    OR COUNT(DISTINCT streams) > 1
                    OR COUNT(DISTINCT track_name) > 1
                    OR COUNT(DISTINCT artist_names) > 1
                    OR COUNT(DISTINCT COALESCE(weekly_start, '')) > 1
                    OR COUNT(DISTINCT COALESCE(weekly_end, '')) > 1
            )
            """,
            (session_id,),
        ).fetchone()
        album_conflicts = int(
            self.connection.execute(
                """
                SELECT COUNT(*) FROM (
                    SELECT album_id
                    FROM staged_chart_rows
                    WHERE session_id=? AND album_id IS NOT NULL
                    GROUP BY album_id
                    HAVING COUNT(DISTINCT album_name) > 1
                        OR COUNT(DISTINCT COALESCE(album_artist, '')) > 1
                        OR COUNT(DISTINCT COALESCE(release_date, '')) > 1
                )
                """,
                (session_id,),
            ).fetchone()[0]
        )
        artist_conflicts = int(
            self.connection.execute(
                """
                SELECT COUNT(*) FROM (
                    SELECT artist_id
                    FROM staged_chart_rows
                    WHERE session_id=? AND artist_id IS NOT NULL
                    GROUP BY artist_id
                    HAVING COUNT(DISTINCT artist_name) > 1
                )
                """,
                (session_id,),
            ).fetchone()[0]
        )
        conflict_count = (
            int(conflict_row["conflicts"]) + album_conflicts + artist_conflicts
        )
        track_counts = self.connection.execute(
            """
            SELECT COUNT(DISTINCT track_id) AS distinct_ids,
                   SUM(CASE WHEN track_id LIKE '%|%' THEN 1 ELSE 0 END) AS fallback_rows
            FROM staged_chart_rows WHERE session_id=?
            """,
            (session_id,),
        ).fetchone()
        bounds = self.connection.execute(
            """
            SELECT MIN(chart_date) AS first_date, MAX(chart_date) AS last_date
            FROM staged_chart_rows WHERE session_id=?
            """,
            (session_id,),
        ).fetchone()
        years = {
            row["year"]: int(row["n"])
            for row in self.connection.execute(
                """
                SELECT substr(chart_date, 1, 4) AS year, COUNT(DISTINCT chart_date) AS n
                FROM staged_chart_rows WHERE session_id=?
                GROUP BY substr(chart_date, 1, 4) ORDER BY year
                """,
                (session_id,),
            ).fetchall()
        }
        artist_mapping_rows = int(
            self.connection.execute(
                """
                SELECT COUNT(*) FROM staged_chart_rows
                WHERE session_id=? AND artist_id IS NOT NULL AND artist_name IS NOT NULL
                """,
                (session_id,),
            ).fetchone()[0]
        )
        album_mapping_rows = int(
            self.connection.execute(
                """
                SELECT COUNT(*) FROM staged_chart_rows
                WHERE session_id=? AND album_id IS NOT NULL AND album_name IS NOT NULL
                """,
                (session_id,),
            ).fetchone()[0]
        )
        missing_dates: list[str] = []
        if bounds["first_date"] and bounds["last_date"] and session["chart_type"] == "daily":
            existing_dates = {
                row["chart_date"]
                for row in self.connection.execute(
                    """
                    SELECT DISTINCT chart_date FROM staged_chart_rows WHERE session_id=?
                    """,
                    (session_id,),
                ).fetchall()
            }
            current = date.fromisoformat(bounds["first_date"])
            last = date.fromisoformat(bounds["last_date"])
            while current <= last:
                candidate = current.isoformat()
                if candidate not in existing_dates:
                    missing_dates.append(candidate)
                current = date.fromordinal(current.toordinal() + 1)
        return ImportSessionSummary(
            session_id=session_id,
            source_name=session["source_name"],
            file_path=session["file_path"],
            status=session["status"],
            current_stage=int(session["current_stage"]),
            processed_rows=int(session["rows_processed"]),
            accepted_rows=accepted,
            rejected_rows=sum(rejection_counts.values()),
            exact_duplicates=exact_duplicates,
            conflicts=conflict_count,
            unique_track_ids=int(track_counts["distinct_ids"] or 0),
            fallback_id_rows=int(track_counts["fallback_rows"] or 0),
            country_counts=country_counts,
            chart_type_counts=chart_counts,
            rejection_counts=rejection_counts,
            first_date=bounds["first_date"],
            last_date=bounds["last_date"],
            coverage_by_year=years,
            artist_mapping_rows=artist_mapping_rows,
            album_mapping_rows=album_mapping_rows,
            missing_date_candidates=missing_dates,
            rights_status=session["rights_status"],
            error=session["error"],
        )

    def finalize(
        self,
        session_id: str,
        database_manager: DatabaseManager,
        progress_callback: Callable[[int, str], None] | None = None,
        cancel_event: object | None = None,
    ) -> dict[str, int | str]:
        session = self.get_session(session_id)
        if session["status"] != "staged":
            raise ValueError("Validate the selected dataset before importing it.")
        source_fingerprint = _fingerprint(Path(session["file_path"]))
        if (
            source_fingerprint["sha256"] != session["file_sha256"]
            or source_fingerprint["size_bytes"] != session["file_size"]
            or source_fingerprint["modified_ns"] != session["file_modified_ns"]
        ):
            raise ValueError("Source file changed after validation; production import was blocked.")
        summary = self.summary(session_id)
        if not summary.accepted_rows:
            raise ValueError("There are no valid USA observations to import.")
        if summary.conflicts:
            raise ValueError(
                f"Production import is blocked by {summary.conflicts:,} conflicting track-date identities."
            )
        if session["chart_type"] == "weekly":
            invalid_week = self.connection.execute(
                """
                SELECT COUNT(*) FROM staged_chart_rows
                WHERE session_id=? AND (weekly_start IS NULL OR weekly_end IS NULL)
                """,
                (session_id,),
            ).fetchone()[0]
            if invalid_week:
                raise ValueError("Weekly imports require explicit start and end dates.")
        if self.database_path.resolve() == Path(database_manager.db_path).resolve():
            raise ValueError("Import staging and production databases must be separate files.")

        source_license = session["source_license"] or "User-selected local data; permissions not independently verified"
        connection = database_manager.connection
        connection.execute("ATTACH DATABASE ? AS import_stage", (str(self.database_path),))
        try:
            inserted = processed = 0
            max_source_row = int(
                self.connection.execute(
                    "SELECT MAX(source_row) FROM staged_chart_rows WHERE session_id=?",
                    (session_id,),
                ).fetchone()[0]
            )
            with database_manager.transaction():
                database_manager.ensure_source(
                    session["source_name"],
                    source_license,
                    (
                        f"Imported file SHA-256 {session['file_sha256']}; "
                        "User-selected local import. Source terms still apply; "
                        "this import grants no commercial or redistribution rights."
                    ),
                )
                batch_id = database_manager.create_batch(
                    session["source_name"],
                    source_license,
                    session["chart_type"],
                    "US",
                )
                conflict = connection.execute(
                    """
                    SELECT 1
                    FROM import_stage.staged_chart_rows s
                    JOIN chart_records c
                      ON c.source=? AND c.country='US' AND c.chart_type=?
                     AND c.chart_date=s.chart_date AND c.track_id=s.track_id
                    WHERE s.session_id=?
                      AND (c.position IS NOT s.position OR c.streams != s.streams
                           OR c.track_name != s.track_name
                           OR c.artist_names != s.artist_names
                           OR c.weekly_start IS NOT s.weekly_start
                           OR c.weekly_end IS NOT s.weekly_end)
                    LIMIT 1
                    """,
                    (session["source_name"], session["chart_type"], session_id),
                ).fetchone()
                if conflict:
                    raise ValueError(
                        "Import conflicts with an existing production observation from this source. "
                        "No production records were committed."
                    )
                artist_conflict = connection.execute(
                    """
                    SELECT 1
                    FROM import_stage.staged_chart_rows s
                    JOIN artists a ON a.artist_id=s.artist_id
                    WHERE s.session_id=? AND s.artist_id IS NOT NULL
                      AND s.artist_name IS NOT NULL
                      AND a.artist_name IS NOT NULL AND a.artist_name!=s.artist_name
                    LIMIT 1
                    """,
                    (session_id,),
                ).fetchone()
                if artist_conflict:
                    raise ValueError(
                        "Supplied artist ID conflicts with existing production metadata."
                    )
                album_conflict = connection.execute(
                    """
                    SELECT 1
                    FROM import_stage.staged_chart_rows s
                    JOIN albums a ON a.album_id=s.album_id
                    WHERE s.session_id=? AND s.album_id IS NOT NULL
                      AND ((a.album_name IS NOT NULL AND a.album_name!=s.album_name)
                        OR (a.album_artist IS NOT NULL AND s.album_artist IS NOT NULL
                            AND a.album_artist!=s.album_artist)
                        OR (a.release_date IS NOT NULL AND s.release_date IS NOT NULL
                            AND a.release_date!=s.release_date))
                    LIMIT 1
                    """,
                    (session_id,),
                ).fetchone()
                if album_conflict:
                    raise ValueError(
                        "Supplied album ID conflicts with existing production metadata."
                    )

                last_row = 1
                while last_row < max_source_row:
                    if cancel_event is not None and cancel_event.is_set():
                        raise InterruptedError("Import cancelled; transaction rolled back.")
                    upper_row = last_row + CHUNK_SIZE
                    before_changes = connection.total_changes
                    connection.execute(
                        """
                        INSERT OR IGNORE INTO chart_records(
                            batch_id, source, source_license, chart_type, country,
                            chart_date, weekly_start, weekly_end, position, track_id,
                            track_name, artist_names, streams, original_url
                        )
                        SELECT DISTINCT ?, ?, ?, ?, 'US', s.chart_date,
                            s.weekly_start, s.weekly_end, s.position, s.track_id,
                            s.track_name, s.artist_names, s.streams, s.original_url
                        FROM import_stage.staged_chart_rows s
                        WHERE s.session_id=? AND s.source_row>? AND s.source_row<=?
                          AND NOT EXISTS (
                              SELECT 1 FROM chart_records c
                              WHERE c.source=? AND c.country='US' AND c.chart_type=?
                                AND c.chart_date=s.chart_date AND c.track_id=s.track_id
                                AND c.position IS s.position AND c.streams=s.streams
                                AND c.track_name=s.track_name
                                AND c.artist_names=s.artist_names
                          )
                        """,
                        (
                            batch_id,
                            session["source_name"],
                            source_license,
                            session["chart_type"],
                            session_id,
                            last_row,
                            upper_row,
                            session["source_name"],
                            session["chart_type"],
                        ),
                    )
                    inserted += connection.total_changes - before_changes

                    connection.execute(
                        """
                        INSERT OR IGNORE INTO tracks(track_id, track_name, artist_names, source)
                        SELECT track_id, MIN(track_name), MIN(artist_names), ?
                        FROM import_stage.staged_chart_rows
                        WHERE session_id=? AND source_row>? AND source_row<=?
                        GROUP BY track_id
                        """,
                        (session["source_name"], session_id, last_row, upper_row),
                    )
                    connection.execute(
                        """
                        INSERT OR IGNORE INTO artists(artist_id, artist_name)
                        SELECT artist_id, MIN(artist_name)
                        FROM import_stage.staged_chart_rows
                        WHERE session_id=? AND source_row>? AND source_row<=?
                          AND artist_id IS NOT NULL AND artist_name IS NOT NULL
                        GROUP BY artist_id
                        """,
                        (session_id, last_row, upper_row),
                    )
                    connection.execute(
                        """
                        INSERT OR IGNORE INTO track_artists(
                            track_id, artist_id, attribution_rule, is_primary,
                            metadata_batch_id, mapping_status
                        )
                        SELECT DISTINCT track_id, artist_id, 'full-credit', 1, ?,
                            'id_match'
                        FROM import_stage.staged_chart_rows
                        WHERE session_id=? AND source_row>? AND source_row<=?
                          AND artist_id IS NOT NULL AND artist_name IS NOT NULL
                        """,
                        (batch_id, session_id, last_row, upper_row),
                    )
                    connection.execute(
                        """
                        INSERT OR IGNORE INTO albums(
                            album_id, album_name, album_artist, release_date,
                            source_identifier, source_name, source_license,
                            metadata_batch_id
                        )
                        SELECT album_id, MIN(album_name), MIN(album_artist),
                            MIN(release_date), MIN(source_reference), ?, ?, ?
                        FROM import_stage.staged_chart_rows
                        WHERE session_id=? AND source_row>? AND source_row<=?
                          AND album_id IS NOT NULL AND album_name IS NOT NULL
                        GROUP BY album_id
                        """,
                        (
                            session["source_name"],
                            source_license,
                            batch_id,
                            session_id,
                            last_row,
                            upper_row,
                        ),
                    )
                    connection.execute(
                        """
                        INSERT OR IGNORE INTO track_albums(
                            track_id, album_id, mapping_confidence,
                            metadata_batch_id, mapping_status
                        )
                        SELECT DISTINCT track_id, album_id, 'source-identified', ?,
                            'id_match'
                        FROM import_stage.staged_chart_rows
                        WHERE session_id=? AND source_row>? AND source_row<=?
                          AND album_id IS NOT NULL AND album_name IS NOT NULL
                        """,
                        (batch_id, session_id, last_row, upper_row),
                    )
                    processed += int(
                        connection.execute(
                            """
                            SELECT COUNT(*) FROM import_stage.staged_chart_rows
                            WHERE session_id=? AND source_row>? AND source_row<=?
                            """,
                            (session_id, last_row, upper_row),
                        ).fetchone()[0]
                    )
                    last_row = upper_row
                    if progress_callback:
                        progress_callback(
                            min(processed, summary.accepted_rows),
                            "Writing validated records transactionally",
                        )

                duplicates_skipped = max(0, summary.accepted_rows - inserted)
                if progress_callback:
                    progress_callback(
                        summary.accepted_rows,
                        "Checking production database integrity",
                    )
                if cancel_event is not None and cancel_event.is_set():
                    raise InterruptedError("Import cancelled; transaction rolled back.")
                integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
                if integrity != "ok":
                    raise sqlite3.DatabaseError(
                        f"Production database integrity check failed: {integrity}"
                    )
        finally:
            connection.execute("DETACH DATABASE import_stage")
        with self.connection:
            self.connection.execute(
                """
                UPDATE import_sessions SET status='imported', current_stage=13,
                    updated_at=datetime('now') WHERE session_id=?
                """,
                (session_id,),
            )
        return {
            "batch_id": batch_id,
            "processed_rows": processed,
            "imported_rows": inserted,
            "duplicates_skipped": duplicates_skipped,
            "database_size_bytes": Path(database_manager.db_path).stat().st_size,
        }

    def rejection_samples(self, session_id: str) -> list[sqlite3.Row]:
        return self.connection.execute(
            """
            SELECT source_row, reason FROM import_rejection_samples
            WHERE session_id=? ORDER BY source_row
            """,
            (session_id,),
        ).fetchall()


def summary_to_dict(summary: ImportSessionSummary) -> dict[str, object]:
    return asdict(summary)
