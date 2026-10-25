from __future__ import annotations

import csv
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .db import DatabaseManager
from .exports import export_rankings
from .metadata_importer import import_metadata_dataframe


def run_packaged_export_checks() -> int:
    try:
        with tempfile.TemporaryDirectory(prefix="usmca-release-export-check-") as directory:
            root = Path(directory)
            db = DatabaseManager(root / "isolated-smoke.sqlite")
            try:
                fixture = [
                    ("A" * 22, "Song Alpha", "Artist Alpha", 1200),
                    ("B" * 22, "Song Beta", "Artist Beta", 800),
                ]
                for index, (track_id, title, artist, streams) in enumerate(fixture, start=1):
                    db.register_chart_record(
                        batch_id="release-smoke",
                        source="INTERNAL RELEASE QA / SYNTHETIC",
                        source_license="Ephemeral smoke fixture; generated only for local verification.",
                        chart_type="daily",
                        country="US",
                        chart_date="2024-01-01",
                        weekly_start=None,
                        weekly_end=None,
                        position=index,
                        track_id=track_id,
                        track_name=title,
                        artist_names=artist,
                        streams=streams,
                        original_url=None,
                    )
                import_metadata_dataframe(
                    db,
                    pd.DataFrame(
                        [
                            {
                                "track_id": "A" * 22,
                                "album_id": "E" * 22,
                                "album_name": "Album Alpha",
                                "artist_id": "C" * 22,
                                "artist_name": "Artist Alpha",
                            },
                            {
                                "track_id": "B" * 22,
                                "album_id": "F" * 22,
                                "album_name": "Album Beta",
                                "artist_id": "D" * 22,
                                "artist_name": "Artist Beta",
                            },
                        ]
                    ),
                    "INTERNAL RELEASE QA / SYNTHETIC",
                    "Ephemeral smoke fixture; generated only for local verification.",
                    "runtime-generated-metadata.csv",
                    {
                        "track_id": "track_id",
                        "album_id": "album_id",
                        "album_name": "album_name",
                        "artist_id": "artist_id",
                        "artist_name": "artist_name",
                    },
                )

                for category in ("songs", "artists", "albums"):
                    for extension in (".csv", ".xlsx"):
                        destination = root / f"{category}{extension}"
                        exported = export_rankings(
                            db,
                            category,
                            destination,
                            "2024-01-01",
                            "2024-01-01",
                            "daily",
                            50,
                        )
                        if exported != 2 or not destination.is_file() or destination.stat().st_size == 0:
                            raise RuntimeError(f"{category} {extension} export produced unexpected output.")
                        if extension == ".csv":
                            with destination.open(encoding="utf-8-sig", newline="") as handle:
                                content = list(csv.reader(handle))
                            if not any(row and row[0] == "Rank" for row in content):
                                raise RuntimeError(f"{category} CSV is missing its rankings header.")
                        else:
                            workbook = load_workbook(destination, read_only=True)
                            try:
                                if workbook.sheetnames != ["Rankings", "Definitions", "Coverage & Sources"]:
                                    raise RuntimeError(f"{category} XLSX has unexpected report worksheets.")
                            finally:
                                workbook.close()
            finally:
                db.close()
        if sys.stdout is not None:
            print("PASS: packaged runtime created and verified all six isolated CSV/XLSX exports.")
        return 0
    except Exception as exc:
        if sys.stderr is not None:
            print(f"FAIL: packaged export smoke test failed: {exc}", file=sys.stderr)
        return 1
