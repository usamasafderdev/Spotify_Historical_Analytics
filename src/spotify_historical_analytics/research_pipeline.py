from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import sqlite3
import uuid
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from shutil import disk_usage
from typing import Any, Callable, Iterable, Iterator, Mapping, Sequence
from urllib.parse import urlsplit


UNVERIFIED = "UNVERIFIED"
APPROVED = "APPROVED"
RESTRICTED = "RESTRICTED"
REJECTED = "REJECTED"
RIGHTS_STATES = {UNVERIFIED, APPROVED, RESTRICTED, REJECTED}
STAGING_SCHEMA_VERSION = 4
PROCESSING_VERSION = 3
REJECTION_SAMPLE_LIMIT = 20
RANK_REJECTION_SAMPLE_LIMIT = 100


@dataclass(frozen=True)
class DatasetSpec:
    dataset_id: str
    owner: str
    slug: str
    version: int
    url: str
    publisher: str
    listed_license: str
    research_folder: str

    @property
    def handle(self) -> str:
        return f"{self.owner}/{self.slug}/versions/{self.version}"


DATASETS = (
    DatasetSpec(
        dataset_id="gonzalopezgil/spotify-charts-daily-updated",
        owner="gonzalopezgil",
        slug="spotify-charts-daily-updated",
        version=146,
        url="https://www.kaggle.com/datasets/gonzalopezgil/spotify-charts-daily-updated",
        publisher="Gonzalo López Gil",
        listed_license="CC BY-SA 4.0 (Kaggle listing; upstream rights unverified)",
        research_folder="01_Charts_Daily_Updated",
    ),
    DatasetSpec(
        dataset_id="dhruvildave/spotify-charts",
        owner="dhruvildave",
        slug="spotify-charts",
        version=85,
        url="https://www.kaggle.com/datasets/dhruvildave/spotify-charts",
        publisher="Dhruvil Dave",
        listed_license="ODbL 1.0 database / original-author contents (Kaggle listing)",
        research_folder="02_Dhruvil_Spotify_Charts",
    ),
)

ALIASES: dict[str, tuple[str, ...]] = {
    "chart_date": ("date", "chart_date", "chartdate", "period", "week", "weekly_start"),
    "weekly_start": ("weekly_start", "week_start", "weekstart"),
    "weekly_end": ("weekly_end", "week_end", "weekend"),
    "country": ("region", "country", "market", "territory", "country_code", "region_code", "market_code"),
    "chart": ("chart", "chart_type", "chart_name", "category"),
    "position": ("rank", "position", "chart_position"),
    "track_id": ("track_id", "spotify_track_id", "uri"),
    "track_url": ("spotify_track_url", "track_url", "url"),
    "track_name": ("track_name", "track", "title", "song", "song_name"),
    "artist_id": ("artist_id", "spotify_artist_id"),
    "artist_name": ("artist_name", "artist", "artists", "credited_artist"),
    "album_id": ("album_id", "spotify_album_id"),
    "album_name": ("album_name", "album", "album_title"),
    "streams": ("streams", "stream", "stream_count", "total_streams"),
}


def _column_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.strip().lower())


def detect_columns(headers: Iterable[str]) -> dict[str, str]:
    by_key: dict[str, str] = {}
    for header in headers:
        by_key.setdefault(_column_key(str(header)), str(header))
    result: dict[str, str] = {}
    for canonical, candidates in ALIASES.items():
        for candidate in candidates:
            if _column_key(candidate) in by_key:
                result[canonical] = by_key[_column_key(candidate)]
                break
    return result


def normalize_region(value: Any) -> str | None:
    if value is None:
        return None
    normalized = re.sub(r"[^a-z]", "", str(value).strip().casefold())
    if normalized in {"us", "usa", "unitedstates", "unitedstatesofamerica"}:
        return "US"
    return None


def classify_category(file_name: str, columns: Mapping[str, str]) -> str:
    name = Path(file_name).name.casefold()
    if "album" in name or "album_name" in columns:
        return "albums"
    if "artist" in name and "song" not in name and "track" not in name:
        return "artists"
    if any(token in name for token in ("song", "track")) or "track_name" in columns:
        return "songs"
    return "unknown"


def _coerce_date(value: Any) -> str | None:
    if value is None or not str(value).strip():
        return None
    text = str(value).strip()
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        pass
    for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _coerce_integer(value: Any) -> int | None:
    if value is None or not str(value).strip():
        return None
    text = str(value).strip().replace(",", "")
    if text.casefold() in {"null", "none", "na", "n/a", "nan"}:
        return None
    try:
        number = float(text)
    except (ValueError, TypeError):
        raise ValueError("not_numeric") from None
    if not math.isfinite(number) or not number.is_integer():
        raise ValueError("not_integer")
    result = int(number)
    if result < 0:
        raise ValueError("negative_value")
    return result


def _value(row: Mapping[str, Any], columns: Mapping[str, str], field: str) -> Any:
    column = columns.get(field)
    return row.get(column) if column else None


def _canonical_record(
    row: Mapping[str, Any],
    columns: Mapping[str, str],
    category: str,
    file_name: str,
) -> tuple[dict[str, Any] | None, str | None]:
    chart_date = _coerce_date(_value(row, columns, "chart_date"))
    if chart_date is None:
        return None, "invalid_or_missing_chart_date"
    country = normalize_region(_value(row, columns, "country"))
    if country is None:
        return None, "non_us_or_unknown_region"
    raw_position = _value(row, columns, "position")
    try:
        position = _coerce_integer(raw_position)
    except ValueError as exc:
        return None, f"invalid_position:{exc}"
    if position is None or position == 0:
        return None, "missing_or_invalid_position"
    chart_name = _value(row, columns, "chart")
    chart_name = str(chart_name).strip() if chart_name is not None else ""
    rank_limit = 50 if "viral" in chart_name.casefold() else 200
    if position > rank_limit:
        return None, f"position_out_of_range:{rank_limit}"

    raw_streams = _value(row, columns, "streams")
    try:
        streams = _coerce_integer(raw_streams)
    except ValueError as exc:
        return None, f"invalid_stream_count:{exc}"

    file_lower = file_name.casefold()
    cadence = "weekly" if "weekly" in file_lower or "weekly" in chart_name.casefold() else "daily"
    weekly_start = _coerce_date(_value(row, columns, "weekly_start"))
    weekly_end = _coerce_date(_value(row, columns, "weekly_end"))
    track_name = _value(row, columns, "track_name")
    artist_name = _value(row, columns, "artist_name")
    album_name = _value(row, columns, "album_name")
    record = {
        "category": category,
        "cadence": cadence,
        "chart_name": chart_name or None,
        "country": country,
        "chart_date": chart_date,
        "weekly_start": weekly_start,
        "weekly_end": weekly_end,
        "position": position,
        "track_id": (
            _clean_text(_value(row, columns, "track_id"))
            or _track_id_from_url(_value(row, columns, "track_url"))
        ),
        "track_name": _clean_text(track_name),
        "artist_id": _clean_text(_value(row, columns, "artist_id")),
        "artist_name": _clean_text(artist_name),
        "album_id": _clean_text(_value(row, columns, "album_id")),
        "album_name": _clean_text(album_name),
        "streams": streams,
    }
    if category == "songs" and record["track_name"] is None:
        return None, "missing_track_name"
    if category == "artists" and record["artist_name"] is None:
        return None, "missing_artist_name"
    if category == "albums" and record["album_name"] is None:
        return None, "missing_album_name"
    return record, None


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _track_id_from_url(value: Any) -> str | None:
    text = _clean_text(value)
    if text is None:
        return None
    parsed = urlsplit(text)
    segments = parsed.path.strip("/").split("/")
    if (
        parsed.scheme.casefold() != "https"
        or (parsed.hostname or "").casefold() != "open.spotify.com"
        or len(segments) != 2
        or segments[0] != "track"
        or re.fullmatch(r"[A-Za-z0-9]{22}", segments[1]) is None
    ):
        return None
    return segments[1]


def sha256_file(path: Path, chunk_bytes: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(chunk_bytes):
            digest.update(chunk)
    return digest.hexdigest()


def file_format(path: Path) -> str:
    name = path.name.casefold()
    if name.endswith(".csv.gz"):
        return "csv.gz"
    return path.suffix.casefold().lstrip(".") or "unknown"


def iter_table_batches(
    path: Path,
    chunk_size: int,
) -> Iterator[tuple[list[str], list[dict[str, Any]]]]:
    fmt = file_format(path)
    if fmt in {"csv", "csv.gz"}:
        import gzip

        opener = gzip.open if fmt == "csv.gz" else open
        with opener(path, "rt", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            csv_headers = list(reader.fieldnames or [])
            csv_batch: list[dict[str, Any]] = []
            for row in reader:
                csv_batch.append(row)
                if len(csv_batch) >= chunk_size:
                    yield csv_headers, csv_batch
                    csv_batch = []
            if csv_batch:
                yield csv_headers, csv_batch
        return
    if fmt == "parquet":
        import pyarrow.parquet as parquet

        table = parquet.ParquetFile(path)
        for record_batch in table.iter_batches(batch_size=chunk_size):
            yield list(record_batch.schema.names), record_batch.to_pylist()
        return
    if fmt in {"xlsx", "xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header_values = next(rows, None)
                if not header_values:
                    continue
                xlsx_headers = [str(value or "") for value in header_values]
                xlsx_batch: list[dict[str, Any]] = []
                for values in rows:
                    xlsx_batch.append(dict(zip(xlsx_headers, values)))
                    if len(xlsx_batch) >= chunk_size:
                        yield xlsx_headers, xlsx_batch
                        xlsx_batch = []
                if xlsx_batch:
                    yield xlsx_headers, xlsx_batch
        finally:
            workbook.close()
        return
    if fmt in {"jsonl", "ndjson"}:
        jsonl_batch: list[dict[str, Any]] = []
        jsonl_headers: list[str] = []
        with path.open("r", encoding="utf-8-sig") as stream:
            for line in stream:
                if not line.strip():
                    continue
                item = json.loads(line)
                if not isinstance(item, dict):
                    raise ValueError("JSONL rows must be JSON objects.")
                if not jsonl_headers:
                    jsonl_headers = list(item)
                jsonl_batch.append(item)
                if len(jsonl_batch) >= chunk_size:
                    yield jsonl_headers, jsonl_batch
                    jsonl_batch = []
        if jsonl_batch:
            yield jsonl_headers, jsonl_batch
        return
    raise ValueError(f"Unsupported row-oriented file format: {fmt}")


def _initialize_staging(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        PRAGMA foreign_keys = ON;
        CREATE TABLE IF NOT EXISTS staged_files (
            dataset_id TEXT NOT NULL,
            dataset_version INTEGER NOT NULL,
            file_sha256 TEXT NOT NULL,
            file_path TEXT NOT NULL,
            status TEXT NOT NULL,
            processing_version INTEGER NOT NULL DEFAULT 1,
            rows_seen INTEGER NOT NULL DEFAULT 0,
            rows_accepted INTEGER NOT NULL DEFAULT 0,
            rows_rejected INTEGER NOT NULL DEFAULT 0,
            started_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            completed_at TEXT,
            error TEXT,
            PRIMARY KEY (dataset_id, dataset_version, file_sha256)
        );
        CREATE TABLE IF NOT EXISTS staged_observations (
            id INTEGER PRIMARY KEY,
            dataset_id TEXT NOT NULL,
            dataset_version INTEGER NOT NULL,
            file_sha256 TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            category TEXT NOT NULL,
            cadence TEXT NOT NULL,
            chart_name TEXT,
            country TEXT NOT NULL,
            chart_date TEXT NOT NULL,
            position INTEGER NOT NULL,
            weekly_start TEXT,
            weekly_end TEXT,
            track_id TEXT,
            track_name TEXT,
            artist_id TEXT,
            artist_name TEXT,
            album_id TEXT,
            album_name TEXT,
            streams INTEGER,
            raw_json TEXT NOT NULL,
            UNIQUE(dataset_id, dataset_version, file_sha256, source_row)
        );
        CREATE TABLE IF NOT EXISTS staged_rejections (
            id INTEGER PRIMARY KEY,
            dataset_id TEXT NOT NULL,
            dataset_version INTEGER NOT NULL,
            file_sha256 TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            reason TEXT NOT NULL,
            raw_json TEXT NOT NULL,
            UNIQUE(dataset_id, dataset_version, file_sha256, source_row)
        );
        CREATE TABLE IF NOT EXISTS staged_rejection_counts (
            dataset_id TEXT NOT NULL,
            dataset_version INTEGER NOT NULL,
            file_sha256 TEXT NOT NULL,
            reason TEXT NOT NULL,
            detail TEXT NOT NULL DEFAULT '',
            rejected_count INTEGER NOT NULL,
            PRIMARY KEY(dataset_id, dataset_version, file_sha256, reason, detail)
        );
        CREATE TABLE IF NOT EXISTS staged_rejection_samples (
            dataset_id TEXT NOT NULL,
            dataset_version INTEGER NOT NULL,
            file_sha256 TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            reason TEXT NOT NULL,
            detail TEXT NOT NULL DEFAULT '',
            raw_json TEXT NOT NULL,
            PRIMARY KEY(dataset_id, dataset_version, file_sha256, source_row)
        );
        CREATE TABLE IF NOT EXISTS source_rights (
            dataset_id TEXT NOT NULL,
            dataset_version INTEGER NOT NULL,
            rights_status TEXT NOT NULL DEFAULT 'UNVERIFIED',
            evidence_path TEXT,
            reviewed_by TEXT,
            reviewed_at TEXT,
            notes TEXT,
            CHECK (rights_status IN ('UNVERIFIED', 'APPROVED', 'RESTRICTED', 'REJECTED')),
            PRIMARY KEY (dataset_id, dataset_version)
        );
        CREATE INDEX IF NOT EXISTS idx_staged_observations_date
            ON staged_observations(category, cadence, chart_date);
        CREATE INDEX IF NOT EXISTS idx_staged_observations_identity
            ON staged_observations(category, cadence, chart_name, chart_date, track_id, position);
        """
    )
    existing_columns = {
        row[1] for row in connection.execute("PRAGMA table_info(staged_observations)")
    }
    for column, declaration in (
        ("weekly_start", "TEXT"),
        ("weekly_end", "TEXT"),
        ("raw_json", "TEXT NOT NULL DEFAULT '{}'"),
    ):
        if column not in existing_columns:
            connection.execute(
                f"ALTER TABLE staged_observations ADD COLUMN {column} {declaration}"
            )
    staged_file_columns = {
        row[1] for row in connection.execute("PRAGMA table_info(staged_files)")
    }
    if "processing_version" not in staged_file_columns:
        connection.execute(
            "ALTER TABLE staged_files ADD COLUMN processing_version INTEGER NOT NULL DEFAULT 1"
        )
    rights_columns = {
        row[1] for row in connection.execute("PRAGMA table_info(source_rights)")
    }
    if "dataset_version" not in rights_columns:
        connection.execute("ALTER TABLE source_rights RENAME TO source_rights_legacy")
        connection.execute(
            """
            CREATE TABLE source_rights (
                dataset_id TEXT NOT NULL,
                dataset_version INTEGER NOT NULL,
                rights_status TEXT NOT NULL DEFAULT 'UNVERIFIED',
                evidence_path TEXT,
                reviewed_by TEXT,
                reviewed_at TEXT,
                notes TEXT,
                CHECK (rights_status IN ('UNVERIFIED', 'APPROVED', 'RESTRICTED', 'REJECTED')),
                PRIMARY KEY (dataset_id, dataset_version)
            )
            """
        )
        connection.execute(
            """
            INSERT INTO source_rights(
                dataset_id, dataset_version, rights_status, evidence_path,
                reviewed_by, reviewed_at, notes
            )
            SELECT dataset_id, 0, rights_status, evidence_path, reviewed_by, reviewed_at, notes
            FROM source_rights_legacy
            """
        )
        connection.execute("DROP TABLE source_rights_legacy")
    connection.execute(f"PRAGMA user_version = {STAGING_SCHEMA_VERSION}")
    connection.commit()


def process_file(
    connection: sqlite3.Connection,
    spec: DatasetSpec,
    path: Path,
    *,
    chunk_size: int = 50_000,
    file_sha256: str | None = None,
    progress_callback: Callable[[Mapping[str, Any]], None] | None = None,
    force_reprocess: bool = False,
) -> dict[str, int | str]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero")
    digest = file_sha256 or sha256_file(path)
    key = (spec.dataset_id, spec.version, digest)
    current = connection.execute(
        "SELECT status, rows_seen, rows_accepted, rows_rejected, processing_version "
        "FROM staged_files "
        "WHERE dataset_id=? AND dataset_version=? AND file_sha256=?",
        key,
    ).fetchone()
    if current and current[0] == "complete":
        if current[4] == PROCESSING_VERSION:
            return {
                "status": "reused",
                "rows_seen": current[1],
                "rows_accepted": current[2],
                "rows_rejected": current[3],
            }
        if not force_reprocess:
            return {
                "status": "reprocess_required",
                "rows_seen": current[1],
                "rows_accepted": current[2],
                "rows_rejected": current[3],
            }

    connection.execute("BEGIN IMMEDIATE")
    connection.execute(
        "DELETE FROM staged_observations WHERE dataset_id=? AND dataset_version=? AND file_sha256=?",
        key,
    )
    connection.execute(
        "DELETE FROM staged_rejections WHERE dataset_id=? AND dataset_version=? AND file_sha256=?",
        key,
    )
    connection.execute(
        "DELETE FROM staged_rejection_counts WHERE dataset_id=? AND dataset_version=? AND file_sha256=?",
        key,
    )
    connection.execute(
        "DELETE FROM staged_rejection_samples WHERE dataset_id=? AND dataset_version=? AND file_sha256=?",
        key,
    )
    connection.execute(
        """
        INSERT INTO staged_files
            (dataset_id, dataset_version, file_sha256, file_path, status, processing_version, rows_seen,
             rows_accepted, rows_rejected, started_at, completed_at, error)
        VALUES (?, ?, ?, ?, 'processing', ?, 0, 0, 0, CURRENT_TIMESTAMP, NULL, NULL)
        ON CONFLICT(dataset_id, dataset_version, file_sha256) DO UPDATE SET
            file_path=excluded.file_path, status='processing',
            processing_version=excluded.processing_version, rows_seen=0,
            rows_accepted=0, rows_rejected=0, started_at=CURRENT_TIMESTAMP,
            completed_at=NULL, error=NULL
        """,
        (*key, str(path), PROCESSING_VERSION),
    )
    connection.commit()

    rows_seen = rows_accepted = rows_rejected = source_row = 0
    sample_counts: Counter[str] = Counter()
    try:
        category: str | None = None
        columns: dict[str, str] | None = None
        for headers, batch in iter_table_batches(path, chunk_size):
            if columns is None:
                columns = detect_columns(headers)
                category = classify_category(path.name, columns)
            assert category is not None
            accepted: list[tuple[Any, ...]] = []
            rejection_counts: Counter[tuple[str, str]] = Counter()
            rejection_samples: list[tuple[Any, ...]] = []
            for raw in batch:
                source_row += 1
                rows_seen += 1
                record, reason = _canonical_record(raw, columns, category, path.name)
                if reason:
                    detail = _rejection_detail(raw, columns, reason)
                    rejection_counts[(reason, detail)] += 1
                    sample_limit = (
                        RANK_REJECTION_SAMPLE_LIMIT
                        if reason.startswith("position_out_of_range:")
                        else REJECTION_SAMPLE_LIMIT
                    )
                    if sample_counts[reason] < sample_limit:
                        rejection_samples.append(
                            (
                                *key,
                                source_row,
                                reason,
                                detail,
                                json.dumps(raw, default=str, ensure_ascii=True),
                            )
                        )
                        sample_counts[reason] += 1
                    rows_rejected += 1
                else:
                    assert record is not None
                    accepted.append(
                        (
                            *key,
                            source_row,
                            record["category"],
                            record["cadence"],
                            record["chart_name"],
                            record["country"],
                            record["chart_date"],
                            record["position"],
                            record["weekly_start"],
                            record["weekly_end"],
                            record["track_id"],
                            record["track_name"],
                            record["artist_id"],
                            record["artist_name"],
                            record["album_id"],
                            record["album_name"],
                            record["streams"],
                            json.dumps(raw, default=str, ensure_ascii=True),
                        )
                    )
                    rows_accepted += 1

            connection.execute("BEGIN IMMEDIATE")
            connection.executemany(
                """
                INSERT INTO staged_observations (
                    dataset_id, dataset_version, file_sha256, source_row, category,
                    cadence, chart_name, country, chart_date, position, weekly_start, weekly_end, track_id,
                    track_name, artist_id, artist_name, album_id, album_name, streams, raw_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                accepted,
            )
            connection.executemany(
                """
                INSERT INTO staged_rejection_counts (
                    dataset_id, dataset_version, file_sha256, reason, detail, rejected_count
                ) VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(dataset_id, dataset_version, file_sha256, reason, detail)
                DO UPDATE SET rejected_count=staged_rejection_counts.rejected_count
                    + excluded.rejected_count
                """,
                [
                    (*key, reason, detail, count)
                    for (reason, detail), count in rejection_counts.items()
                ],
            )
            connection.executemany(
                """
                INSERT OR IGNORE INTO staged_rejection_samples (
                    dataset_id, dataset_version, file_sha256, source_row,
                    reason, detail, raw_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                rejection_samples,
            )
            connection.execute(
                """
                UPDATE staged_files SET rows_seen=?, rows_accepted=?, rows_rejected=?
                WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
                """,
                (rows_seen, rows_accepted, rows_rejected, *key),
            )
            connection.commit()
            if progress_callback is not None:
                progress_callback(
                    {
                        "stage": "processing",
                        "rows_seen": rows_seen,
                        "rows_accepted": rows_accepted,
                        "rows_rejected": rows_rejected,
                    }
                )

        if columns is None:
            raise ValueError("No table rows or headers were found.")
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            """
            UPDATE staged_files SET status='complete', completed_at=CURRENT_TIMESTAMP
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            """,
            key,
        )
        connection.execute(
            "INSERT OR IGNORE INTO source_rights(dataset_id, dataset_version, rights_status) "
            "VALUES (?, ?, 'UNVERIFIED')",
            (spec.dataset_id, spec.version),
        )
        connection.commit()
        if progress_callback is not None:
            progress_callback(
                {
                    "stage": "complete",
                    "rows_seen": rows_seen,
                    "rows_accepted": rows_accepted,
                    "rows_rejected": rows_rejected,
                }
            )
    except Exception as exc:
        connection.rollback()
        connection.execute(
            """
            UPDATE staged_files SET status='failed', error=?, completed_at=CURRENT_TIMESTAMP
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            """,
            (f"{type(exc).__name__}: {exc}", *key),
        )
        connection.commit()
        raise

    return {
        "status": "complete",
        "rows_seen": rows_seen,
        "rows_accepted": rows_accepted,
        "rows_rejected": rows_rejected,
    }


def _rejection_detail(
    row: Mapping[str, Any],
    columns: Mapping[str, str],
    reason: str,
) -> str:
    if reason == "non_us_or_unknown_region":
        raw_value = _clean_text(_value(row, columns, "country"))
        return "missing" if raw_value is None else re.sub(r"\s+", " ", raw_value)[:100]
    if reason.startswith("position_out_of_range:"):
        chart = _clean_text(_value(row, columns, "chart")) or "unknown"
        rank = _clean_text(_value(row, columns, "position")) or "missing"
        return f"{chart[:60]}:rank={rank[:20]}"
    return ""


def supplement_rank_rejection_samples(
    source_database: Path,
    source_file: Path,
    compact_database: Path,
    spec: DatasetSpec,
) -> int:
    digest = sha256_file(source_file)
    source_connection = sqlite3.connect(
        f"{source_database.resolve().as_uri()}?mode=ro", uri=True
    )
    compact_connection = sqlite3.connect(compact_database)
    try:
        source_file_row = source_connection.execute(
            """
            SELECT rows_rejected FROM staged_files
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=? AND status='complete'
            """,
            (spec.dataset_id, spec.version, digest),
        ).fetchone()
        compact_file_row = compact_connection.execute(
            """
            SELECT 1 FROM staged_files
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=? AND status='complete'
            """,
            (spec.dataset_id, spec.version, digest),
        ).fetchone()
        if source_file_row is None or compact_file_row is None:
            raise ValueError("Both databases must contain the same completed source checksum.")
        source_tables = {
            row[0]
            for row in source_connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )
        }
        if "staged_rejections" not in source_tables:
            return 0
        headers, _ = _schema_for(source_file)
        columns = detect_columns(headers)
        rows = source_connection.execute(
            """
            SELECT source_row, reason, raw_json
            FROM staged_rejections
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
              AND reason LIKE 'position_out_of_range:%'
            ORDER BY source_row LIMIT ?
            """,
            (
                spec.dataset_id,
                spec.version,
                digest,
                RANK_REJECTION_SAMPLE_LIMIT,
            ),
        ).fetchall()
        samples = []
        for source_row, reason, raw_json in rows:
            raw = json.loads(raw_json)
            detail = _rejection_detail(raw, columns, reason)
            samples.append(
                (
                    spec.dataset_id,
                    spec.version,
                    digest,
                    source_row,
                    reason,
                    detail,
                    raw_json,
                )
            )
        before = compact_connection.total_changes
        compact_connection.executemany(
            """
            INSERT OR IGNORE INTO staged_rejection_samples (
                dataset_id, dataset_version, file_sha256, source_row,
                reason, detail, raw_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            samples,
        )
        compact_connection.commit()
        added = compact_connection.total_changes - before
        integrity = compact_connection.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise sqlite3.DatabaseError(f"Compact staging integrity check failed: {integrity}")
        return added
    finally:
        source_connection.close()
        compact_connection.close()


def build_compact_staging_database(
    source_database: Path,
    source_file: Path,
    target_database: Path,
    spec: DatasetSpec,
    *,
    chunk_size: int = 50_000,
    minimum_free_bytes: int | None = None,
    progress_callback: Callable[[Mapping[str, Any]], None] | None = None,
) -> dict[str, Any]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero")
    source_database = source_database.resolve()
    source_file = source_file.resolve()
    target_database = target_database.resolve()
    target_database.parent.mkdir(parents=True, exist_ok=True)
    temporary_database = target_database.with_name(target_database.name + ".building")
    if target_database.exists():
        raise FileExistsError(f"Refusing to overwrite existing compact database: {target_database}")
    if temporary_database.exists():
        raise FileExistsError(
            f"Refusing to overwrite an existing build file: {temporary_database}"
        )
    if not source_database.is_file() or not source_file.is_file():
        raise FileNotFoundError("The source database and source data file must both exist.")

    digest = sha256_file(source_file)
    source_connection = sqlite3.connect(f"{source_database.as_uri()}?mode=ro", uri=True)
    target_connection: sqlite3.Connection | None = None
    try:
        source_file_row = source_connection.execute(
            """
            SELECT status, rows_seen, rows_accepted, rows_rejected
            FROM staged_files
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            """,
            (spec.dataset_id, spec.version, digest),
        ).fetchone()
        if source_file_row is None:
            raise ValueError("The source database has no matching file checksum and dataset version.")
        status, rows_seen, rows_accepted, rows_rejected = source_file_row
        if status != "complete":
            raise ValueError(f"Source staging status must be complete, not {status!r}.")

        accepted_raw_bytes = source_connection.execute(
            """
            SELECT COALESCE(SUM(length(raw_json)), 0)
            FROM staged_observations
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            """,
            (spec.dataset_id, spec.version, digest),
        ).fetchone()[0]
        required_free_bytes = max(
            int(minimum_free_bytes or 0),
            int(accepted_raw_bytes) + int(rows_accepted) * 2048 + 512 * 1024 * 1024,
        )
        free_bytes = disk_usage(target_database.parent).free
        if free_bytes < required_free_bytes:
            raise OSError(
                f"Insufficient free space for safe compaction: {free_bytes} bytes available, "
                f"{required_free_bytes} bytes required by the conservative estimate."
            )

        target_connection = sqlite3.connect(temporary_database)
        _initialize_staging(target_connection)
        if progress_callback is not None:
            progress_callback(
                {
                    "stage": "preflight_complete",
                    "free_bytes": free_bytes,
                    "required_free_bytes": required_free_bytes,
                    "rows_expected": rows_seen,
                }
            )
        result = process_file(
            target_connection,
            spec,
            source_file,
            chunk_size=chunk_size,
            file_sha256=digest,
            progress_callback=progress_callback,
        )
        if result["status"] != "complete":
            raise RuntimeError(f"Compact staging did not complete: {result['status']}")
        rank_samples_added = supplement_rank_rejection_samples(
            source_database,
            source_file,
            temporary_database,
            spec,
        )
        if progress_callback is not None:
            progress_callback(
                {
                    "stage": "rank_samples_supplemented",
                    "added_samples": rank_samples_added,
                }
            )

        compact_counts = target_connection.execute(
            """
            SELECT rows_seen, rows_accepted, rows_rejected
            FROM staged_files
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            """,
            (spec.dataset_id, spec.version, digest),
        ).fetchone()
        rejection_count = target_connection.execute(
            """
            SELECT COALESCE(SUM(rejected_count), 0)
            FROM staged_rejection_counts
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            """,
            (spec.dataset_id, spec.version, digest),
        ).fetchone()[0]
        observations_count = target_connection.execute(
            """
            SELECT COUNT(*) FROM staged_observations
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            """,
            (spec.dataset_id, spec.version, digest),
        ).fetchone()[0]
        if compact_counts != (rows_seen, rows_accepted, rows_rejected):
            raise ValueError(
                f"Row reconciliation failed: source={(rows_seen, rows_accepted, rows_rejected)}, "
                f"compact={compact_counts}"
            )
        if rejection_count != rows_rejected or observations_count != rows_accepted:
            raise ValueError(
                "Compact accepted/rejected table counts do not reconcile with the source manifest."
            )

        rights = source_connection.execute(
            """
            SELECT rights_status, evidence_path, reviewed_by, reviewed_at, notes
            FROM source_rights WHERE dataset_id=? AND dataset_version=?
            """,
            (spec.dataset_id, spec.version),
        ).fetchone()
        if rights is not None:
            target_connection.execute(
                """
                UPDATE source_rights
                SET rights_status=?, evidence_path=?, reviewed_by=?, reviewed_at=?, notes=?
                WHERE dataset_id=? AND dataset_version=?
                """,
                (*rights, spec.dataset_id, spec.version),
            )
            target_connection.commit()

        integrity = target_connection.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise sqlite3.DatabaseError(f"Compact staging integrity check failed: {integrity}")
        if sha256_file(source_file) != digest:
            raise ValueError("The source data file changed during compact staging.")
        rejection_samples = target_connection.execute(
            """
            SELECT COUNT(*) FROM staged_rejection_samples
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            """,
            (spec.dataset_id, spec.version, digest),
        ).fetchone()[0]
        target_connection.close()
        target_connection = None
        if target_database.exists():
            raise FileExistsError(
                f"Target appeared during compaction; refusing to overwrite: {target_database}"
            )
        temporary_database.rename(target_database)
        return {
            "status": "complete",
            "source_database": str(source_database),
            "source_file": str(source_file),
            "source_sha256": digest,
            "target_database": str(target_database),
            "rows_seen": rows_seen,
            "rows_accepted": rows_accepted,
            "rows_rejected": rows_rejected,
            "rejection_samples": rejection_samples,
            "rank_rejection_samples_added": rank_samples_added,
            "estimated_required_free_bytes": required_free_bytes,
            "free_bytes_at_preflight": free_bytes,
            "integrity_check": integrity,
        }
    finally:
        if target_connection is not None:
            target_connection.close()
        source_connection.close()


def approve_source_rights(connection: sqlite3.Connection, approval_path: Path) -> None:
    payload = json.loads(approval_path.read_text(encoding="utf-8"))
    required = {
        "dataset_id",
        "dataset_version",
        "reviewed_by",
        "reviewed_at",
        "evidence_path",
        "rights_scopes",
    }
    missing = required - payload.keys()
    if missing:
        raise ValueError(f"Approval evidence is missing fields: {', '.join(sorted(missing))}")
    if not str(payload["reviewed_by"]).strip() or not str(payload["evidence_path"]).strip():
        raise ValueError("Approval requires a reviewer and evidence reference.")
    rights_scopes = payload["rights_scopes"]
    required_scopes = {
        "acquisition",
        "derived_analytics",
        "commercial_use",
        "redistribution",
    }
    if not isinstance(rights_scopes, dict) or not required_scopes.issubset(rights_scopes):
        raise ValueError(
            "Approval must document acquisition, derived_analytics, commercial_use, "
            "and redistribution rights scopes."
        )
    if any(rights_scopes[scope] is not True for scope in required_scopes):
        raise PermissionError(
            "Production import requires documented permission for acquisition, "
            "derived analytics, commercial use, and redistribution."
        )
    evidence = Path(payload["evidence_path"])
    if not evidence.is_absolute():
        evidence = (approval_path.parent / evidence).resolve()
    if not evidence.exists():
        raise ValueError("The cited rights evidence file does not exist.")
    if not evidence.is_file():
        raise ValueError("The cited rights evidence path must reference a file.")
    dataset_id = str(payload["dataset_id"])
    version = int(payload["dataset_version"])
    if connection.execute(
        "SELECT 1 FROM staged_files WHERE dataset_id=? AND dataset_version=? AND status='complete' LIMIT 1",
        (dataset_id, version),
    ).fetchone() is None:
        raise ValueError("No fully staged files exist for this source and version.")
    connection.execute(
        """
        INSERT INTO source_rights(
            dataset_id, dataset_version, rights_status, evidence_path, reviewed_by, reviewed_at, notes
        )
        VALUES (?, ?, 'APPROVED', ?, ?, ?, ?)
        ON CONFLICT(dataset_id, dataset_version) DO UPDATE SET
            rights_status='APPROVED', evidence_path=excluded.evidence_path,
            reviewed_by=excluded.reviewed_by, reviewed_at=excluded.reviewed_at,
            notes=excluded.notes
        """,
        (
            dataset_id,
            version,
            str(evidence),
            str(payload["reviewed_by"]),
            str(payload["reviewed_at"]),
            json.dumps(
                {
                    "reviewer_notes": str(payload.get("notes", "")),
                    "rights_scopes": {
                        scope: True for scope in sorted(required_scopes)
                    },
                    "evidence_sha256": sha256_file(evidence),
                    "review_means_evidence_recorded_not_independently_verified": True,
                },
                sort_keys=True,
            ),
        ),
    )
    connection.commit()


def _schema_for(path: Path) -> tuple[list[str], int | None]:
    fmt = file_format(path)
    if fmt in {"csv", "csv.gz"}:
        import gzip

        opener = gzip.open if fmt == "csv.gz" else open
        with opener(path, "rt", encoding="utf-8-sig", newline="") as stream:
            header = next(csv.reader(stream), [])
        return header, None
    if fmt == "parquet":
        import pyarrow.parquet as parquet

        table = parquet.ParquetFile(path)
        return list(table.schema.names), table.metadata.num_rows
    if fmt in {"xlsx", "xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            headers: list[str] = []
            row_count = 0
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                sheet_headers = [str(value or "") for value in next(rows, ())]
                if not headers:
                    headers = sheet_headers
                row_count += max(sheet.max_row - 1, 0)
            return headers, row_count
        finally:
            workbook.close()
    if fmt in {"jsonl", "ndjson"}:
        headers: list[str] = []
        row_count = 0
        with path.open("r", encoding="utf-8-sig") as stream:
            for line in stream:
                if not line.strip():
                    continue
                row_count += 1
                if not headers:
                    item = json.loads(line)
                    headers = list(item) if isinstance(item, dict) else []
        return headers, row_count
    return [], None


def inventory_file(
    path: Path,
    dataset_id: str,
    version: int,
    source_url: str,
    license_text: str,
    *,
    digest: str | None = None,
    known_row_count: int | None = None,
) -> dict[str, Any]:
    size = path.stat().st_size
    headers, known_rows = _schema_for(path) if file_format(path) in {
        "csv",
        "csv.gz",
        "parquet",
        "xlsx",
        "xlsm",
        "jsonl",
        "ndjson",
    } else ([], None)
    row_count = known_row_count if known_row_count is not None else known_rows
    if row_count is None and headers and file_format(path) in {"csv", "csv.gz"}:
        row_count = 0
        for _, batch in iter_table_batches(path, 100_000):
            row_count += len(batch)
    return {
        "dataset_id": dataset_id,
        "source_url": source_url,
        "dataset_version": version,
        "filename": path.name,
        "path": str(path),
        "format": file_format(path),
        "size_bytes": size,
        "sha256": digest or sha256_file(path),
        "row_count": row_count,
        "column_names_json": json.dumps(headers, ensure_ascii=True),
        "detected_fields_json": json.dumps(detect_columns(headers), ensure_ascii=True),
        "category": classify_category(path.name, detect_columns(headers)),
        "listed_license": license_text,
        "permission_status": UNVERIFIED,
    }


def discover_dataset_files(root: Path) -> list[Path]:
    if not root.exists():
        return []
    excluded = {".partial", ".part", ".tmp", ".download", ".archive", ".complete"}
    result: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.casefold() in excluded:
            continue
        if any(part.casefold() in {"audit", "reports", "staging", ".locks"} for part in path.parts):
            continue
        result.append(path)
    return sorted(result)


def _zip_members(path: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise zipfile.BadZipFile(f"CRC failed for archive member {bad_member}")
        return [
            {
                "name": item.filename,
                "size_bytes": item.file_size,
                "compressed_bytes": item.compress_size,
                "format": Path(item.filename).suffix.casefold().lstrip("."),
                "is_directory": item.is_dir(),
            }
            for item in archive.infolist()
            if not item.is_dir()
        ]


def write_inventory_report(
    research_root: Path,
    dataset_roots: Mapping[str, Path],
    report_dir: Path,
) -> list[dict[str, Any]]:
    report_dir.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []
    by_id = {item.dataset_id: item for item in DATASETS}
    for dataset_id, root in dataset_roots.items():
        spec = by_id[dataset_id]
        files = discover_dataset_files(root)
        if not files:
            records.append(
                {
                    "dataset_id": dataset_id,
                    "source_url": spec.url,
                    "dataset_version": spec.version,
                    "filename": "",
                    "path": str(root),
                    "format": "",
                    "size_bytes": 0,
                    "sha256": "",
                    "row_count": "",
                    "column_names_json": "[]",
                    "detected_fields_json": "{}",
                    "category": "",
                    "listed_license": spec.listed_license,
                    "permission_status": UNVERIFIED,
                    "status": "NOT_DOWNLOADED",
                }
            )
            continue
        for path in files:
            if path.suffix.casefold() == ".zip":
                archive = _zip_members(path)
                records.append(
                    {
                        **inventory_file(path, dataset_id, spec.version, spec.url, spec.listed_license),
                        "archive_members_json": json.dumps(archive, ensure_ascii=True),
                        "status": "ARCHIVE_INTEGRITY_OK",
                    }
                )
            else:
                records.append(
                    {
                        **inventory_file(path, dataset_id, spec.version, spec.url, spec.listed_license),
                        "archive_members_json": "",
                        "status": "DISCOVERED",
                    }
                )
    fieldnames = sorted({key for row in records for key in row})
    with (report_dir / "dataset_inventory.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)
    manifest = {
        "research_root": str(research_root),
        "generated_at": datetime.now().astimezone().isoformat(),
        "datasets": [
            {
                **asdict(spec),
                "rights_status": UNVERIFIED,
                "download_status": next(
                    (row.get("status") for row in records if row["dataset_id"] == spec.dataset_id),
                    "NOT_DOWNLOADED",
                ),
            }
            for spec in DATASETS
        ],
    }
    (report_dir / "source_registry.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=True), encoding="utf-8"
    )
    return records


def write_analysis_reports(connection: sqlite3.Connection, report_dir: Path) -> dict[str, int]:
    report_dir.mkdir(parents=True, exist_ok=True)
    report_counts = {
        "coverage_by_year": 0,
        "coverage_by_month": 0,
        "missing_dates": 0,
        "source_overlap": 0,
        "data_conflicts": 0,
    }
    source_join = """
        FROM staged_observations o
        JOIN staged_files f
          ON f.dataset_id=o.dataset_id
         AND f.dataset_version=o.dataset_version
         AND f.file_sha256=o.file_sha256
        WHERE f.status='complete'
    """

    def write_csv(
        name: str,
        fields: Sequence[str],
        records: Iterable[Mapping[str, Any]],
    ) -> int:
        path = report_dir / name
        row_count = 0
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            for record in records:
                writer.writerow(record)
                row_count += 1
        if row_count == 0:
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=["status"])
                writer.writeheader()
                writer.writerow({"status": "NO_STAGED_OBSERVATIONS"})
        return row_count

    def duplicate_counts(period_expression: str) -> dict[tuple[Any, ...], int]:
        query = f"""
            SELECT dataset_id, dataset_version, category, cadence, chart_name, period,
                   SUM(source_rows - 1)
            FROM (
                SELECT o.dataset_id, o.dataset_version, o.category, o.cadence,
                       COALESCE(o.chart_name, '') AS chart_name,
                       {period_expression} AS period, o.chart_date, o.track_id,
                       o.position, COUNT(*) AS source_rows
                {source_join}
                  AND o.track_id IS NOT NULL
                GROUP BY o.dataset_id, o.dataset_version, o.category, o.cadence,
                         COALESCE(o.chart_name, ''), period, o.chart_date,
                         o.track_id, o.position
                HAVING COUNT(*) > 1
            )
            GROUP BY dataset_id, dataset_version, category, cadence, chart_name, period
        """
        return {
            tuple(row[:6]): int(row[6])
            for row in connection.execute(query)
        }

    def aggregate_coverage(
        period_expression: str,
        period_kind: str,
        duplicates: Mapping[tuple[Any, ...], int],
    ) -> list[dict[str, Any]]:
        query = f"""
            SELECT o.dataset_id, o.dataset_version, o.category, o.cadence,
                   COALESCE(o.chart_name, '') AS chart_name,
                   {period_expression} AS period,
                   COUNT(*) AS observation_count,
                   COUNT(DISTINCT o.chart_date) AS distinct_dates,
                   COUNT(DISTINCT NULLIF(o.track_id, '')) AS distinct_track_ids,
                   COUNT(DISTINCT COALESCE(NULLIF(o.artist_id, ''), NULLIF(o.artist_name, '')))
                       AS distinct_artists,
                   COUNT(DISTINCT NULLIF(o.album_id, '')) AS distinct_albums,
                   SUM(CASE WHEN o.streams IS NOT NULL THEN 1 ELSE 0 END) AS rows_with_streams,
                   MIN(o.chart_date), MAX(o.chart_date)
            {source_join}
            GROUP BY o.dataset_id, o.dataset_version, o.category, o.cadence,
                     COALESCE(o.chart_name, ''), period
            ORDER BY o.dataset_id, o.category, o.cadence, chart_name, period
        """
        summaries: list[dict[str, Any]] = []
        for row in connection.execute(query):
            dataset_id, version, category, cadence, chart_name, period = row[:6]
            period_text = str(period)
            period_value: int | str = (
                int(period_text) if period_kind == "year" else period_text
            )
            expected_dates = (
                (date(period_value + 1, 1, 1) - date(period_value, 1, 1)).days
                if period_kind == "year" and isinstance(period_value, int)
                else None
            )
            key = (
                dataset_id, version, category, cadence, chart_name, period_text
            )
            summaries.append(
                {
                    "dataset_id": dataset_id,
                    "dataset_version": version,
                    "category": category,
                    "cadence": cadence,
                    "chart_name": chart_name,
                    "period": period_value,
                    "first_chart_date": row[12],
                    "last_chart_date": row[13],
                    "observation_count": row[6],
                    "distinct_dates": row[7],
                    "distinct_track_ids": row[8],
                    "distinct_artists": row[9],
                    "distinct_albums": row[10],
                    "rows_with_streams": row[11],
                    "duplicate_source_rows": duplicates.get(key, 0),
                    "calendar_day_coverage_percent": (
                        round(row[7] * 100 / expected_dates, 4)
                        if expected_dates
                        else ""
                    ),
                    "coverage_assumption": (
                        "Calendar-day diagnostic only; not proof of expected chart schedule"
                    ),
                    "coverage_status": "OBSERVATIONS_FOUND_SCHEDULE_UNVERIFIED",
                    "permission_status": UNVERIFIED,
                }
            )
        return summaries

    year_duplicates = duplicate_counts("substr(o.chart_date, 1, 4)")
    month_duplicates = duplicate_counts("substr(o.chart_date, 1, 7)")
    year_rows = aggregate_coverage(
        "substr(o.chart_date, 1, 4)", "year", year_duplicates
    )
    month_rows = aggregate_coverage(
        "substr(o.chart_date, 1, 7)", "month", month_duplicates
    )
    expected_series = {
        (spec.dataset_id, category, cadence, "")
        for spec, categories in zip(
            DATASETS,
            (
                (("songs", "daily"), ("artists", "daily"), ("albums", "weekly")),
                (("songs", "daily"),),
            ),
        )
        for category, cadence in categories
    }
    observed_years = {
        (
            row["dataset_id"], row["category"], row["cadence"],
            row["chart_name"], row["period"],
        )
        for row in year_rows
    }
    for dataset_id, category, cadence, chart_name in sorted(expected_series):
        for year in range(2017, 2027):
            key = (dataset_id, category, cadence, chart_name, year)
            if key not in observed_years:
                year_rows.append({
                    "dataset_id": dataset_id,
                    "dataset_version": next(
                        spec.version for spec in DATASETS if spec.dataset_id == dataset_id
                    ),
                    "category": category,
                    "cadence": cadence,
                    "chart_name": chart_name,
                    "period": year,
                    "first_chart_date": "",
                    "last_chart_date": "",
                    "observation_count": 0,
                    "distinct_dates": 0,
                    "distinct_track_ids": 0,
                    "distinct_artists": 0,
                    "distinct_albums": 0,
                    "rows_with_streams": 0,
                    "duplicate_source_rows": 0,
                    "calendar_day_coverage_percent": "",
                    "coverage_assumption": (
                        "No observations staged; expected schedule and source availability unverified"
                    ),
                    "coverage_status": "NO_VALIDATED_OBSERVATIONS",
                    "permission_status": UNVERIFIED,
                    "category_basis": (
                        "Dataset listing only; not verified against downloaded file schemas"
                    ),
                })
    year_rows.sort(
        key=lambda item: (
            item["dataset_id"], item["category"], item["cadence"],
            str(item["chart_name"]), int(item["period"]),
        )
    )
    coverage_fields = (
        "dataset_id", "dataset_version", "category", "cadence", "chart_name",
        "period", "first_chart_date", "last_chart_date", "observation_count",
        "distinct_dates", "distinct_track_ids", "distinct_artists",
        "distinct_albums", "rows_with_streams", "duplicate_source_rows",
        "calendar_day_coverage_percent", "coverage_assumption", "coverage_status",
        "permission_status", "category_basis",
    )
    report_counts["coverage_by_year"] = write_csv(
        "coverage_by_year.csv", coverage_fields, year_rows
    )
    report_counts["coverage_by_month"] = write_csv(
        "coverage_by_month.csv", coverage_fields, month_rows
    )

    missing_fields = (
        "dataset_id", "category", "cadence", "chart_name",
        "candidate_missing_chart_date", "interpretation",
    )
    missing_path = report_dir / "missing_dates.csv"
    with missing_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=missing_fields)
        writer.writeheader()
        series_query = f"""
            SELECT o.dataset_id, o.category, o.cadence,
                   COALESCE(o.chart_name, ''), MIN(o.chart_date)
            {source_join}
            GROUP BY o.dataset_id, o.category, o.cadence, COALESCE(o.chart_name, '')
            ORDER BY o.dataset_id, o.category, o.cadence, COALESCE(o.chart_name, '')
        """
        for series in connection.execute(series_query):
            first_date = date.fromisoformat(series[4])
            step = 7 if series[2] == "weekly" else 1
            candidate = first_date
            dates_query = f"""
                SELECT DISTINCT o.chart_date
                {source_join}
                  AND o.dataset_id=? AND o.category=? AND o.cadence=?
                  AND COALESCE(o.chart_name, '')=? ORDER BY o.chart_date
            """
            for observed in connection.execute(
                dates_query, (series[0], series[1], series[2], series[3])
            ):
                observed_date = date.fromisoformat(observed[0])
                while candidate < observed_date:
                    writer.writerow(
                        {
                            "dataset_id": series[0],
                            "category": series[1],
                            "cadence": series[2],
                            "chart_name": series[3],
                            "candidate_missing_chart_date": candidate.isoformat(),
                            "interpretation": (
                                "No observation found on this calendar/weekly-step date; "
                                "not a zero-stream record and expected schedule is unverified"
                            ),
                        }
                    )
                    report_counts["missing_dates"] += 1
                    candidate = date.fromordinal(candidate.toordinal() + step)
                if candidate == observed_date:
                    candidate = date.fromordinal(candidate.toordinal() + step)
    if report_counts["missing_dates"] == 0:
        with missing_path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=["status"])
            writer.writeheader()
            writer.writerow({"status": "NO_STAGED_OBSERVATIONS"})

    overlap_fields = (
        "category", "cadence", "chart_name", "chart_date", "track_id",
        "positions_json", "track_names_json", "credited_artists_json",
        "datasets_json", "observations", "streams_json", "status",
    )
    overlap_path = report_dir / "source_overlap.csv"
    conflict_path = report_dir / "data_conflicts.csv"
    comparison_query = f"""
        SELECT o.category, o.cadence, COALESCE(o.chart_name, ''), o.chart_date,
               o.track_id, COUNT(*), COUNT(DISTINCT o.dataset_id),
               COUNT(DISTINCT o.position),
               COUNT(DISTINCT COALESCE(o.track_name, '')),
               COUNT(DISTINCT COALESCE(o.artist_name, '')),
               COUNT(DISTINCT o.streams),
               SUM(CASE WHEN o.streams IS NULL THEN 1 ELSE 0 END),
               json_group_array(DISTINCT o.position),
               json_group_array(DISTINCT COALESCE(o.track_name, '')),
               json_group_array(DISTINCT COALESCE(o.artist_name, '')),
               json_group_array(DISTINCT o.dataset_id),
               json_group_array(DISTINCT o.streams)
        {source_join}
          AND o.track_id IS NOT NULL
        GROUP BY o.category, o.cadence, COALESCE(o.chart_name, ''),
                 o.chart_date, o.track_id
        HAVING COUNT(DISTINCT o.dataset_id) > 1
        ORDER BY o.category, o.cadence, o.chart_date, o.track_id
    """
    with (
        overlap_path.open("w", encoding="utf-8-sig", newline="") as overlap_stream,
        conflict_path.open("w", encoding="utf-8-sig", newline="") as conflict_stream,
    ):
        overlap_writer = csv.DictWriter(overlap_stream, fieldnames=overlap_fields)
        conflict_writer = csv.DictWriter(conflict_stream, fieldnames=overlap_fields)
        overlap_writer.writeheader()
        conflict_writer.writeheader()
        for row in connection.execute(comparison_query):
            streams = json.loads(row[16])
            streams.sort(key=lambda value: (-1 if value is None else value))
            comparison = {
                "category": row[0],
                "cadence": row[1],
                "chart_name": row[2],
                "chart_date": row[3],
                "track_id": row[4],
                "positions_json": json.dumps(sorted(json.loads(row[12]))),
                "track_names_json": json.dumps(sorted(json.loads(row[13]))),
                "credited_artists_json": json.dumps(sorted(json.loads(row[14]))),
                "datasets_json": json.dumps(sorted(json.loads(row[15]))),
                "observations": row[5],
                "streams_json": json.dumps(streams),
            }
            streams_match = row[10] <= 1 and row[11] in (0, row[5])
            is_match = (
                streams_match and row[7] == 1 and row[8] == 1 and row[9] == 1
            )
            if is_match:
                overlap_writer.writerow({
                    "category": comparison["category"],
                    "cadence": comparison["cadence"],
                    "chart_name": comparison["chart_name"],
                    "chart_date": comparison["chart_date"],
                    "track_id": comparison["track_id"],
                    "positions_json": comparison["positions_json"],
                    "track_names_json": comparison["track_names_json"],
                    "credited_artists_json": comparison["credited_artists_json"],
                    "datasets_json": comparison["datasets_json"],
                    "observations": comparison["observations"],
                    "streams_json": comparison["streams_json"],
                    "status": "MATCHING_DUPLICATE_OBSERVATION",
                })
                report_counts["source_overlap"] += 1
            else:
                conflict_writer.writerow({
                    "category": comparison["category"],
                    "cadence": comparison["cadence"],
                    "chart_name": comparison["chart_name"],
                    "chart_date": comparison["chart_date"],
                    "track_id": comparison["track_id"],
                    "positions_json": comparison["positions_json"],
                    "track_names_json": comparison["track_names_json"],
                    "credited_artists_json": comparison["credited_artists_json"],
                    "datasets_json": comparison["datasets_json"],
                    "observations": comparison["observations"],
                    "streams_json": comparison["streams_json"],
                    "status": "CONFLICTING_OBSERVATION",
                })
                report_counts["data_conflicts"] += 1
    for path, count in (
        (overlap_path, report_counts["source_overlap"]),
        (conflict_path, report_counts["data_conflicts"]),
    ):
        if count == 0:
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=["status"])
                writer.writeheader()
                writer.writerow({"status": "NO_STAGED_OBSERVATIONS"})

    with (report_dir / "deduplication_summary.md").open("w", encoding="utf-8") as stream:
        stream.write("# Source overlap and deduplication summary\n\n")
        stream.write(f"- Matching cross-source duplicate identities: {report_counts['source_overlap']}\n")
        stream.write(f"- Conflicting cross-source identities: {report_counts['data_conflicts']}\n")
        stream.write("- No records were summed, deleted, or selected as a preferred source.\n")
        stream.write("- Comparison requires a stable track ID and matches category, cadence, chart, date, track ID, and position.\n")
        stream.write("- Name-only matches are not treated as verified duplicates.\n")
    with (report_dir / "historical_coverage.md").open("w", encoding="utf-8") as stream:
        stream.write("# Historical chart coverage audit\n\n")
        stream.write("Required client period: **2017–2026**. A calendar-day percentage is diagnostic only; it is not proof that each date was an expected reporting date.\n\n")
        observed_rows = [row for row in year_rows if row["observation_count"]]
        if not observed_rows:
            stream.write("No validated chart observations are staged. No first/last date, coverage year, or USA count can currently be confirmed.\n\n")
        else:
            stream.write("| Dataset | Category | Cadence | Year | Rows | Distinct dates | First | Last | Streams rows |\n")
            stream.write("|---|---|---|---:|---:|---:|---|---|---:|\n")
            for row in observed_rows:
                stream.write(
                    f"| {row['dataset_id']} | {row['category']} | {row['cadence']} | "
                    f"{row['period']} | {row['observation_count']} | {row['distinct_dates']} | "
                    f"{row['first_chart_date']} | {row['last_chart_date']} | {row['rows_with_streams']} |\n"
                )
            stream.write("\n")
        stream.write("Every source remains `UNVERIFIED` for commercial analytics and redistribution. Missing dates are not treated as zero streams. Completeness is not asserted without a verified expected chart calendar.\n")
    aggregated_rejections = connection.execute(
        "SELECT COALESCE(SUM(rejected_count), 0) FROM staged_rejection_counts"
    ).fetchone()[0]
    legacy_rejections = connection.execute(
        "SELECT COUNT(*) FROM staged_rejections"
    ).fetchone()[0]
    validation = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "staged_observations": connection.execute(
            f"SELECT COUNT(*) {source_join}"
        ).fetchone()[0],
        "staged_files": connection.execute("SELECT COUNT(*) FROM staged_files WHERE status='complete'").fetchone()[0],
        "staged_rejections": aggregated_rejections or legacy_rejections,
        "permissions_approved": connection.execute(
            "SELECT COUNT(*) FROM source_rights WHERE rights_status='APPROVED'"
        ).fetchone()[0],
        "production_imports_performed": 0,
        "reports": report_counts,
    }
    (report_dir / "import_validation.json").write_text(
        json.dumps(validation, indent=2), encoding="utf-8"
    )
    return report_counts


def production_import_records(
    connection: sqlite3.Connection,
    database_manager: Any,
    dataset_id: str,
    dataset_version: int,
) -> int:
    rights = connection.execute(
        "SELECT rights_status FROM source_rights WHERE dataset_id=? AND dataset_version=?",
        (dataset_id, dataset_version),
    ).fetchone()
    if rights is None or rights[0] != APPROVED:
        raise PermissionError("Production import is blocked until explicit rights approval evidence is recorded.")
    ready_files = connection.execute(
        """
        SELECT COUNT(*) FROM staged_files
        WHERE dataset_id=? AND dataset_version=? AND status='complete'
        """,
        (dataset_id, dataset_version),
    ).fetchone()[0]
    if not ready_files:
        raise ValueError("No completed staged files exist for this dataset version.")

    sources = connection.execute(
        """
        SELECT DISTINCT category, cadence, chart_name
        FROM staged_observations o
        JOIN staged_files f USING (dataset_id, dataset_version, file_sha256)
        WHERE dataset_id=? AND dataset_version=? AND category='songs' AND f.status='complete'
        """,
        (dataset_id, dataset_version),
    ).fetchall()
    if any(row[0] != "songs" or row[1] not in {"daily", "weekly"} for row in sources):
        raise ValueError(
            "Production import supports only song-chart rows with explicit daily/weekly cadence; "
            "artist/album metrics and mixed categories remain in research staging."
        )
    if any(row[2] and "viral" in str(row[2]).casefold() for row in sources):
        raise ValueError("Viral chart records are not compatible with the current production chart model.")

    source_name = f"{dataset_id}@{dataset_version}"
    app_connection = database_manager.connection
    source_license = f"Approved research source: {dataset_id}@{dataset_version}"
    try:
        app_connection.execute("BEGIN IMMEDIATE")
        app_connection.execute(
            "INSERT OR IGNORE INTO sources(name, license, remarks) VALUES (?, ?, ?)",
            (source_name, source_license, "Imported from a rights-approved research staging database."),
        )
        inserted = 0
        for cadence in ("daily", "weekly"):
            records = connection.execute(
                """
                SELECT chart_date, position, weekly_start, weekly_end, track_id, track_name, artist_name, streams
                FROM staged_observations o
                JOIN staged_files f USING (dataset_id, dataset_version, file_sha256)
                WHERE dataset_id=? AND dataset_version=? AND cadence=? AND category='songs'
                  AND streams IS NOT NULL AND f.status='complete'
                ORDER BY chart_date, position
                """,
                (dataset_id, dataset_version, cadence),
            ).fetchall()
            if not records:
                continue
            if cadence == "weekly" and any(row[2] is None or row[3] is None for row in records):
                raise ValueError(
                    "Weekly song records without explicit week_start and week_end stay in research staging."
                )
            batch_id = uuid.uuid4().hex
            app_connection.execute(
                """
                INSERT INTO import_batches
                    (batch_id, source_name, source_license, chart_type, country, is_demo)
                VALUES (?, ?, ?, ?, 'US', 0)
                """,
                (batch_id, source_name, source_license, cadence),
            )
            before = app_connection.total_changes
            app_connection.executemany(
                """
                INSERT OR IGNORE INTO chart_records (
                    batch_id, source, source_license, chart_type, country, chart_date,
                    weekly_start, weekly_end, position, track_id, track_name,
                    artist_names, streams, original_url
                ) VALUES (?, ?, ?, ?, 'US', ?, ?, ?, ?, ?, ?, ?, ?, NULL)
                """,
                [
                    (
                        batch_id, source_name, source_license, cadence, row[0], row[2],
                        row[3], row[1], row[4], row[5], row[6] or "", row[7],
                    )
                    for row in records
                ],
            )
            inserted += app_connection.total_changes - before
        app_connection.commit()
        return inserted
    except Exception:
        app_connection.rollback()
        raise


def write_chart_quality_reports(
    connection: sqlite3.Connection,
    report_dir: Path,
    dataset_id: str,
    dataset_version: int,
    *,
    start_year: int = 2017,
    end_year: int = 2026,
) -> dict[str, Any]:
    if start_year > end_year:
        raise ValueError("start_year must not exceed end_year")
    report_dir.mkdir(parents=True, exist_ok=True)
    source = connection.execute(
        """
        SELECT file_sha256, rows_seen, rows_accepted, rows_rejected, status
        FROM staged_files
        WHERE dataset_id=? AND dataset_version=?
        ORDER BY completed_at DESC LIMIT 1
        """,
        (dataset_id, dataset_version),
    ).fetchone()
    if source is None or source[4] != "complete":
        raise ValueError("Quality reports require a completed staged dataset file.")
    file_sha256, rows_seen, rows_accepted, rows_rejected, _ = source
    base_where = """
        FROM staged_observations o
        JOIN staged_files f USING (dataset_id, dataset_version, file_sha256)
        WHERE f.status='complete' AND o.dataset_id=? AND o.dataset_version=?
          AND o.file_sha256=? AND o.category='songs' AND o.cadence='daily'
    """
    observations_by_date: dict[str, dict[str, Any]] = {}
    for row in connection.execute(
        f"""
        SELECT chart_name, chart_date, COUNT(*) AS row_count,
               SUM(CASE WHEN streams IS NOT NULL THEN 1 ELSE 0 END) AS stream_rows,
               SUM(CASE WHEN streams=0 THEN 1 ELSE 0 END) AS zero_stream_rows
        {base_where}
        GROUP BY chart_name, chart_date
        ORDER BY chart_name, chart_date
        """,
        (dataset_id, dataset_version, file_sha256),
    ):
        chart_name, chart_date, row_count, stream_rows, zero_stream_rows = row
        key = str(chart_name or "unknown")
        record = observations_by_date.setdefault(
            key,
            {
                "dates": {},
                "first_date": str(chart_date),
                "last_date": str(chart_date),
            },
        )
        record["dates"][str(chart_date)] = {
            "row_count": int(row_count),
            "stream_rows": int(stream_rows or 0),
            "zero_stream_rows": int(zero_stream_rows or 0),
        }
        record["first_date"] = min(record["first_date"], str(chart_date))
        record["last_date"] = max(record["last_date"], str(chart_date))

    chart_names = sorted(set(observations_by_date) | {"top200", "viral50"})
    expected_sizes = {"top200": 200, "viral50": 50}
    year_track_counts = {
        (str(chart_name or "unknown"), str(year)): int(track_count)
        for chart_name, year, track_count in connection.execute(
            f"""
            SELECT o.chart_name, substr(o.chart_date, 1, 4),
                   COUNT(DISTINCT NULLIF(o.track_id, ''))
            {base_where}
            GROUP BY o.chart_name, substr(o.chart_date, 1, 4)
            """,
            (dataset_id, dataset_version, file_sha256),
        )
    }
    duplicate_stats = {
        (str(chart_name or "unknown"), str(year)): {
            "duplicate_groups": int(groups),
            "duplicate_extra_rows": int(extra_rows),
            "conflicting_groups": int(conflicts),
        }
        for chart_name, year, groups, extra_rows, conflicts in connection.execute(
            f"""
            SELECT chart_name, year, COUNT(*), SUM(row_count - 1),
                   SUM(CASE WHEN stream_variants > 1 THEN 1 ELSE 0 END)
            FROM (
                SELECT o.chart_name, substr(o.chart_date, 1, 4) AS year,
                       o.chart_date, o.track_id, o.position, COUNT(*) AS row_count,
                       COUNT(DISTINCT COALESCE(CAST(o.streams AS TEXT), 'NULL'))
                           AS stream_variants
                {base_where} AND o.track_id IS NOT NULL
                GROUP BY o.chart_name, substr(o.chart_date, 1, 4),
                         o.chart_date, o.track_id, o.position
                HAVING COUNT(*) > 1
            )
            GROUP BY chart_name, year
            """,
            (dataset_id, dataset_version, file_sha256),
        )
    }

    daily_fields = (
        "chart_name", "chart_date", "observed_rows", "expected_rank_count",
        "valid_stream_rows", "zero_stream_rows", "status", "calendar_assumption",
    )
    year_fields = (
        "chart_name", "year", "calendar_days", "observed_dates",
        "calendar_missing_dates", "observations", "valid_stream_observations",
        "missing_stream_observations", "zero_stream_observations",
        "incomplete_observed_dates", "expected_rank_count",
        "duplicate_identity_groups", "duplicate_extra_rows",
        "conflicting_stream_identity_groups", "unique_track_ids",
        "first_observed_date", "last_observed_date", "calendar_assumption",
    )
    month_fields = (
        "chart_name", "month", "calendar_days", "observed_dates",
        "calendar_missing_dates", "observations", "valid_stream_observations",
        "incomplete_observed_dates", "expected_rank_count", "calendar_assumption",
    )
    daily_records: list[dict[str, Any]] = []
    year_records: list[dict[str, Any]] = []
    month_totals: dict[tuple[str, str], dict[str, int]] = {}
    first_dates: dict[str, date] = {}
    last_dates: dict[str, date] = {}
    for chart_name in chart_names:
        observed = observations_by_date.get(chart_name, {"dates": {}})
        dates = observed["dates"]
        actual_dates = [date.fromisoformat(value) for value in dates]
        if actual_dates:
            first_dates[chart_name] = min(actual_dates)
            last_dates[chart_name] = max(actual_dates)
        expected = expected_sizes.get(chart_name)
        for year in range(start_year, end_year + 1):
            year_dates = [
                (date(year, 1, 1).toordinal() + offset)
                for offset in range((date(year + 1, 1, 1) - date(year, 1, 1)).days)
            ]
            year_days = [date.fromordinal(value) for value in year_dates]
            observed_year = {
                day.isoformat(): dates[day.isoformat()]
                for day in year_days
                if day.isoformat() in dates
            }
            incomplete = sum(
                1
                for item in observed_year.values()
                if expected is not None and item["row_count"] != expected
            )
            year_observations = sum(item["row_count"] for item in observed_year.values())
            year_streams = sum(item["stream_rows"] for item in observed_year.values())
            year_zeros = sum(item["zero_stream_rows"] for item in observed_year.values())
            span_start = first_dates.get(chart_name)
            span_end = last_dates.get(chart_name)
            year_record = {
                "chart_name": chart_name,
                "year": year,
                "calendar_days": len(year_days),
                "observed_dates": len(observed_year),
                "calendar_missing_dates": len(year_days) - len(observed_year),
                "observations": year_observations,
                "valid_stream_observations": year_streams,
                "missing_stream_observations": year_observations - year_streams,
                "zero_stream_observations": year_zeros,
                "incomplete_observed_dates": incomplete,
                "expected_rank_count": expected or "",
                "duplicate_identity_groups": duplicate_stats.get(
                    (chart_name, str(year)), {}
                ).get("duplicate_groups", 0),
                "duplicate_extra_rows": duplicate_stats.get(
                    (chart_name, str(year)), {}
                ).get("duplicate_extra_rows", 0),
                "conflicting_stream_identity_groups": duplicate_stats.get(
                    (chart_name, str(year)), {}
                ).get("conflicting_groups", 0),
                "unique_track_ids": year_track_counts.get((chart_name, str(year)), 0),
                "first_observed_date": min(observed_year) if observed_year else "",
                "last_observed_date": max(observed_year) if observed_year else "",
                "calendar_assumption": (
                    "Calendar-day comparison only; expected publication schedule is unverified."
                ),
            }
            year_records.append(year_record)
            for month in range(1, 13):
                month_key = f"{year:04d}-{month:02d}"
                start = date(year, month, 1)
                end = date(year + (1 if month == 12 else 0), month % 12 + 1, 1)
                month_days = (end - start).days
                totals = month_totals.setdefault(
                    (chart_name, month_key),
                    {
                        "calendar_days": month_days,
                        "observed_dates": 0,
                        "observations": 0,
                        "valid_stream_observations": 0,
                        "incomplete_observed_dates": 0,
                    },
                )
                for day in year_days:
                    if day.month != month or day.isoformat() not in dates:
                        continue
                    item = dates[day.isoformat()]
                    totals["observed_dates"] += 1
                    totals["observations"] += item["row_count"]
                    totals["valid_stream_observations"] += item["stream_rows"]
                    if expected is not None and item["row_count"] != expected:
                        totals["incomplete_observed_dates"] += 1

                for day in year_days:
                    if day.month != month:
                        continue
                    item = dates.get(day.isoformat())
                    is_inside_span = (
                        span_start is not None
                        and span_end is not None
                        and span_start <= day <= span_end
                    )
                    if item is None and is_inside_span:
                        status = "GAP_CANDIDATE_WITHIN_OBSERVED_SPAN"
                    elif item is None:
                        status = "NO_OBSERVATION_OUTSIDE_OBSERVED_SPAN"
                    elif expected is not None and item["row_count"] != expected:
                        status = "OBSERVED_INCOMPLETE_RANK_COUNT"
                    else:
                        status = "OBSERVED"
                    daily_records.append(
                        {
                            "chart_name": chart_name,
                            "chart_date": day.isoformat(),
                            "observed_rows": 0 if item is None else item["row_count"],
                            "expected_rank_count": expected or "",
                            "valid_stream_rows": 0 if item is None else item["stream_rows"],
                            "zero_stream_rows": 0 if item is None else item["zero_stream_rows"],
                            "status": status,
                            "calendar_assumption": (
                                "Daily calendar candidate only; source reporting schedule is unverified."
                            ),
                        }
                    )
    month_records = [
        {
            "chart_name": chart_name,
            "month": month,
            "calendar_days": totals["calendar_days"],
            "observed_dates": totals["observed_dates"],
            "calendar_missing_dates": totals["calendar_days"] - totals["observed_dates"],
            "observations": totals["observations"],
            "valid_stream_observations": totals["valid_stream_observations"],
            "incomplete_observed_dates": totals["incomplete_observed_dates"],
            "expected_rank_count": expected_sizes.get(chart_name, ""),
            "calendar_assumption": (
                "Calendar-day comparison only; expected publication schedule is unverified."
            ),
        }
        for (chart_name, month), totals in sorted(month_totals.items())
    ]

    def write_csv_file(
        name: str,
        fields: Sequence[str],
        records: Iterable[Mapping[str, Any]],
    ) -> None:
        with (report_dir / name).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(records)

    write_csv_file("chart_quality_by_date.csv", daily_fields, daily_records)
    write_csv_file("chart_quality_by_year.csv", year_fields, year_records)
    write_csv_file("chart_quality_by_month.csv", month_fields, month_records)

    rejection_records = [
        {
            "reason": reason,
            "detail": detail,
            "rejected_rows": count,
        }
        for reason, detail, count in connection.execute(
            """
            SELECT reason, detail, SUM(rejected_count)
            FROM staged_rejection_counts
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
            GROUP BY reason, detail ORDER BY reason, detail
            """,
            (dataset_id, dataset_version, file_sha256),
        )
    ]
    write_csv_file(
        "rejection_summary.csv",
        ("reason", "detail", "rejected_rows"),
        rejection_records,
    )
    sample_records = [
        {
            "source_row": source_row,
            "reason": reason,
            "detail": detail,
            "raw_json": raw_json,
        }
        for source_row, reason, detail, raw_json in connection.execute(
            """
            SELECT source_row, reason, detail, raw_json
            FROM staged_rejection_samples
            WHERE dataset_id=? AND dataset_version=? AND file_sha256=?
              AND reason LIKE 'position_out_of_range:%'
            ORDER BY source_row
            """,
            (dataset_id, dataset_version, file_sha256),
        )
    ]
    write_csv_file(
        "invalid_rank_samples.csv",
        ("source_row", "reason", "detail", "raw_json"),
        sample_records,
    )

    rejection_total = sum(int(row["rejected_rows"]) for row in rejection_records)
    summary = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "dataset_id": dataset_id,
        "dataset_version": dataset_version,
        "file_sha256": file_sha256,
        "source_rows_seen": rows_seen,
        "accepted_rows": rows_accepted,
        "rejected_rows": rows_rejected,
        "rejection_aggregates_reconcile": rejection_total == rows_rejected,
        "accepted_daily_song_observations": sum(
            int(row["observations"]) for row in year_records
        ),
        "chart_summaries": {
            chart_name: {
                "first_observed_date": (
                    first_dates[chart_name].isoformat()
                    if chart_name in first_dates
                    else None
                ),
                "last_observed_date": (
                    last_dates[chart_name].isoformat()
                    if chart_name in last_dates
                    else None
                ),
                "years": [
                    row for row in year_records if row["chart_name"] == chart_name
                ],
            }
            for chart_name in chart_names
        },
        "invalid_rank_sample_count": len(sample_records),
        "invalid_rank_sample_limit_per_reason": RANK_REJECTION_SAMPLE_LIMIT,
        "general_rejection_sample_limit_per_reason": REJECTION_SAMPLE_LIMIT,
        "calendar_coverage_note": (
            "Calendar dates without rows are candidate gaps only. No expected source "
            "publication schedule was verified; missing rows are not zero streams."
        ),
        "permissions_note": (
            "Research staging only. The Kaggle listing license does not by itself "
            "resolve underlying Spotify rights or commercial-use permission."
        ),
    }
    (report_dir / "chart_quality_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=True), encoding="utf-8"
    )
    return summary
