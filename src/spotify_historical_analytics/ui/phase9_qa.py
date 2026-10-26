"""Opt-in Windows/widget acceptance runner. Never opens the user's database."""
from __future__ import annotations

import csv
import json
from pathlib import Path
import time

from PySide6.QtCore import QDate, QTimer, Qt
from PySide6.QtWidgets import QApplication, QDialog, QFileDialog, QLineEdit, QMessageBox, QScrollArea
from openpyxl import load_workbook

from ..db import DatabaseManager
from ..settings import SettingsStore
from .main_window import MainWindow
from .import_dialog import ImportDialog


def run_phase9_qa(output: Path, width: int = 1200, height: int = 760) -> int:
    output.mkdir(parents=True, exist_ok=True)
    fixture = output / f"fixture-{time.time_ns()}"
    fixture.mkdir()
    production = fixture / "database" / "analytics.db"
    database = DatabaseManager(production)
    app = QApplication.instance()
    window = MainWindow(db=database)
    window.resize(width, height)
    window.show()
    captures, checks, errors = [], [], []

    def wait(predicate):
        deadline = time.monotonic() + 30
        while not predicate() and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.005)
        assert predicate(), "Phase 9 GUI check timed out"

    def settled():
        return window._has_current_dashboard_payload() and not window._query_threads

    def capture(widget, name, desktop=False):
        app.processEvents()
        path = output / f"{name}.png"
        assert widget.grab().save(str(path))
        captures.append(path.name)
        if desktop and app.platformName() == "windows":
            from .windows_capture import save_native_window
            native = output / f"{name}-desktop.png"
            if save_native_window(widget, native):
                captures.append(native.name)

    # Close informational messages raised by actual export/backup operations.
    messages = QTimer(window)
    messages.setInterval(50)
    def dismiss_messages():
        for widget in app.topLevelWidgets():
            if isinstance(widget, QMessageBox) and widget.isVisible():
                if widget.icon() == QMessageBox.Critical:
                    errors.append(widget.text())
                if widget.icon() == QMessageBox.Warning:
                    capture(widget, "invalid-date-error", desktop=True)
                widget.accept()
    messages.timeout.connect(dismiss_messages)
    messages.start()

    def select_save_file(path):
        def select():
            for chooser in app.topLevelWidgets():
                if isinstance(chooser, QFileDialog) and chooser.isVisible():
                    chooser.setDirectory(str(path.parent.resolve()))
                    def finish(chooser=chooser):
                        chooser.findChild(QLineEdit, "fileNameEdit").setText(path.name)
                        chooser.accept()
                    QTimer.singleShot(100, finish)
        QTimer.singleShot(100, select)

    wait(settled)
    assert database.connection.execute("SELECT COUNT(*) FROM chart_records").fetchone()[0] == 0
    for page in range(7):
        window.navigate(page)
        capture(window, f"empty-{page + 1:02d}", desktop=page == 0)
    checks.append("All seven pages and first-run empty states")
    window.navigate(0)
    window.empty_demo_button.click()
    wait(settled)
    assert window._demo_mode and window.demo_banner.isVisible()
    assert "DEMO / SYNTHETIC" in window.windowTitle()
    assert window.songs_table.rowCount() == 6
    assert window.artist_table.rowCount() > 0 and window.album_table.rowCount() > 0
    assert window.streams_chart.plot.listDataItems()
    assert window.top_songs_chart.plot.listDataItems()
    scroll = window.findChild(QScrollArea, "contentScroll")
    for page in range(7):
        window.navigate(page)
        scroll.verticalScrollBar().setValue(0)
        capture(window, f"demo-{page + 1:02d}", desktop=page in (0, 1, 2, 3, 4, 6))
        scroll.verticalScrollBar().setValue(scroll.verticalScrollBar().maximum())
        capture(window, f"demo-{page + 1:02d}-content")
        assert scroll.horizontalScrollBar().maximum() == 0
    checks.append("Isolated synthetic preview, populated KPIs/tables/charts, scrolling")

    window.navigate(1)
    window.song_search.setText("no matching song 72931")
    wait(lambda: not window._search_timer.isActive() and settled())
    assert window.songs_table.rowCount() == 0 and window.songs_empty.isVisible()
    capture(window, "empty-search")
    window.song_search.setText("Demo Song 1")
    wait(lambda: not window._search_timer.isActive() and settled())
    assert window.songs_table.rowCount() == 1
    window.song_search.clear()
    wait(lambda: not window._search_timer.isActive() and settled())
    window.songs_table.sortItems(3, Qt.DescendingOrder)
    values = [window.songs_table.item(row, 3).data(Qt.UserRole) for row in range(window.songs_table.rowCount())]
    assert values == sorted(values, reverse=True)
    checks.append("Song search, empty search, numeric sorting")

    # Open each actual asynchronous table detail action and dismiss its modal UI.
    for kind, page, table in [("songs", 1, window.songs_table), ("artists", 2, window.artist_table), ("albums", 3, window.album_table)]:
        opened = []
        timer = QTimer(window)
        timer.setInterval(30)
        def close_detail(kind=kind):
            for dialog in app.topLevelWidgets():
                if isinstance(dialog, QDialog) and not isinstance(dialog, QMessageBox) and dialog.isVisible():
                    capture(dialog, f"detail-{kind}", desktop=True)
                    opened.append(True)
                    dialog.accept()
        timer.timeout.connect(close_detail)
        timer.start()
        window.navigate(page)
        table.setCurrentCell(0, 1)
        getattr(window, {"songs": "_open_song_detail_from_selection", "artists": "_open_artist_detail_from_selection", "albums": "_open_album_detail_from_selection"}[kind])()
        wait(lambda: bool(opened) and not window._detail_threads)
        timer.stop()
        timer.deleteLater()
    checks.append("Song, artist and album detail dialogs and trend charts")

    window.start_edit.setDate(QDate(2020, 1, 1))
    window.end_edit.setDate(QDate(2020, 12, 31))
    window.apply_filters()
    wait(settled)
    assert window.active_start == "2020-01-01" and window.active_end == "2020-12-31"
    exports = fixture / "exports"
    exports.mkdir()
    for cadence in ("daily", "weekly"):
        window.chart_type_combo.setCurrentIndex(window.chart_type_combo.findData(cadence))
        wait(settled)
        assert window.active_chart_type == cadence
        for category, page, button in [("songs", 1, window.song_export_button), ("artists", 2, window.artist_export_button), ("albums", 3, window.album_export_button)]:
            window.navigate(page)
            assert button.isEnabled()
            for extension in (".csv", ".xlsx"):
                path = exports / f"demo-{cadence}-{category}{extension}"
                select_save_file(path)
                # The actual save dialog chooses CSV/XLSX from the filename suffix.
                button.click()
                wait(lambda: not window._export_threads)
                assert path.is_file() and not errors, errors
                if extension == ".csv":
                    with path.open(encoding="utf-8-sig", newline="") as handle:
                        rows = list(csv.reader(handle))
                    header = next(i for i, row in enumerate(rows) if row and row[0] == "Rank")
                    assert len(rows[header + 1:]) > 0
                    assert any("DEMO / SYNTHETIC" in " ".join(row) for row in rows)
                else:
                    book = load_workbook(path, read_only=True, data_only=True)
                    assert book["Rankings"].max_row > 1
                    assert "DEMO / SYNTHETIC" in str(dict(list(book["Definitions"].values)[1:])["Synthetic data"])
                    book.close()
    checks.append("Actual CSV/XLSX save controls and parsed exports for all rankings, daily/weekly, filtered synthetic data")
    capture(window, "filtered-weekly-demo", desktop=True)
    window.demo_button.click()
    wait(settled)
    assert not window._demo_mode and not window.overview_data.isVisible()
    assert database.connection.execute("SELECT COUNT(*) FROM chart_records").fetchone()[0] == 0
    assert not window.metadata_button.isEnabled()
    checks.append("Return to real data restores empty production without contamination")
    window.navigate(6)
    backup = fixture / "backup.db"
    select_save_file(backup)
    window.backup_button.click()
    wait(lambda: backup.is_file())
    check_db = DatabaseManager(backup)
    assert check_db.connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    check_db.close()
    settings = SettingsStore()
    assert settings.get("last_export_directory") == str(exports.resolve())
    checks.append("Settings persistence and actual backup save control")
    dialog = ImportDialog(production, fixture / "staging.db", window)
    dialog.show()
    capture(dialog, "simple-import", desktop=True)
    dialog.close()
    window.start_edit.setDate(QDate(2021, 1, 2))
    window.end_edit.setDate(QDate(2020, 1, 1))
    QTimer.singleShot(150, lambda: [capture(w, "invalid-date-error", desktop=True) for w in app.topLevelWidgets() if isinstance(w, QMessageBox) and w.isVisible()])
    window.apply_filters()
    wait(settled)
    capture(window, "return-to-empty", desktop=True)
    messages.stop()
    scale = window.devicePixelRatioF()
    window.close()
    reopened = DatabaseManager(production)
    assert reopened.connection.execute("SELECT COUNT(*) FROM chart_records").fetchone()[0] == 0
    reopened.close()
    checks.append("Database restart persistence; production remains empty")
    (output / "results.json").write_text(json.dumps({
        "checks": checks, "captures": captures, "platform": app.platformName(),
        "device_pixel_ratio": scale, "logical_size": [width, height],
        "production_records": 0, "exported_reports": 12, "errors": errors,
    }, indent=2))
    return 0
