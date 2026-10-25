from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from spotify_historical_analytics.db import DatabaseManager
from spotify_historical_analytics.exports import export_rankings
from spotify_historical_analytics.metadata_importer import import_metadata_dataframe


TRACK_A = "A" * 22
TRACK_B = "B" * 22
ARTIST_A = "C" * 22
ARTIST_B = "D" * 22
ALBUM_A = "E" * 22
ALBUM_B = "F" * 22


def create_export_database(path: Path) -> DatabaseManager:
    db = DatabaseManager(path)
    for chart_date, track_id, name, artist, streams in [
        ("2024-01-01", TRACK_A, "Song Alpha", "Credited Alpha", 100),
        ("2024-01-02", TRACK_A, "Song Alpha", "Credited Alpha", 300),
        ("2024-01-01", TRACK_B, "Song Beta", "Credited Beta", 200),
    ]:
        db.register_chart_record(
            batch_id="export-fixture",
            source="Authorized export fixture",
            source_license="Fixture license note",
            chart_type="daily",
            country="US",
            chart_date=chart_date,
            weekly_start=None,
            weekly_end=None,
            position=1,
            track_id=track_id,
            track_name=name,
            artist_names=artist,
            streams=streams,
            original_url=None,
        )
    metadata = pd.DataFrame(
        [
            {
                "track_id": TRACK_A,
                "album_id": ALBUM_A,
                "album_name": "Album Alpha",
                "artist_id": ARTIST_A,
                "artist_name": "Artist Alpha",
            },
            {
                "track_id": TRACK_B,
                "album_id": ALBUM_B,
                "album_name": "Album Beta",
                "artist_id": ARTIST_B,
                "artist_name": "Artist Beta",
            },
        ]
    )
    import_metadata_dataframe(
        db,
        metadata,
        "Authorized export metadata",
        "Client authorized fixture",
        "export-metadata.csv",
        {
            "track_id": "track_id",
            "album_id": "album_id",
            "album_name": "album_name",
            "artist_id": "artist_id",
            "artist_name": "artist_name",
        },
    )
    return db


@pytest.mark.parametrize(
    ("category", "search", "expected_value"),
    [
        ("songs", "Song Alpha", "Total chart streams"),
        ("artists", "Artist Alpha", "Charting-song streams"),
        ("albums", "Album Alpha", "Album charting-track streams"),
    ],
)
@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_filtered_export_matches_database_rankings(
    tmp_path: Path, category: str, search: str, expected_value: str, suffix: str
) -> None:
    db = create_export_database(tmp_path / "export.db")
    destination = tmp_path / f"{category}{suffix}"
    if category == "songs":
        expected_total = db.get_top_tracks(
            50, "2024-01-01", "2024-01-01", chart_type="daily", search=search
        )[0]["total_streams"]
        metric_column = 3
    elif category == "artists":
        expected_total = db.get_artist_summary(
            "2024-01-01", "2024-01-01", chart_type="daily", limit=50, search=search
        )[0]["charting_song_streams"]
        metric_column = 3
    else:
        expected_total = db.get_album_rankings(
            "2024-01-01", "2024-01-01", "daily", limit=50, search=search
        )[0]["charting_track_streams"]
        metric_column = 4
    count = export_rankings(
        db,
        category,
        destination,
        "2024-01-01",
        "2024-01-01",
        "daily",
        50,
        search,
    )
    assert count == 1
    assert destination.exists()

    if suffix == ".csv":
        with destination.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))
        assert ["Market", "United States (US)"] in rows
        assert ["Chart type", "daily"] in rows
        assert ["Source names", "Authorized export fixture"] in rows
        assert ["Observed chart dates", "1"] in rows
        ranking_header_index = rows.index(
            next(row for row in rows if row and row[0] == "Rank")
        )
        ranking_headers = rows[ranking_header_index]
        data_row = rows[ranking_header_index + 1]
        assert ranking_headers[3 if category == "songs" else 3 if category == "artists" else 4] == expected_value
        assert data_row[0] == "1"
        assert int(data_row[metric_column]) == expected_total
    else:
        workbook = load_workbook(destination, read_only=True, data_only=True)
        assert workbook.sheetnames == ["Rankings", "Definitions", "Coverage & Sources"]
        rankings = workbook["Rankings"]
        assert rankings.max_row == 2
        assert expected_value in [cell.value for cell in rankings[1]]
        assert rankings.cell(2, metric_column + 1).value == expected_total
        assert workbook["Definitions"]["B2"].value == category.title()
        source_rows = list(workbook["Coverage & Sources"].values)
        assert any("Authorized export fixture" in row for row in source_rows)
        assert any("Fixture license note" in row for row in source_rows)


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_empty_export_contains_definitions_and_no_ranking_rows(
    tmp_path: Path, suffix: str
) -> None:
    db = create_export_database(tmp_path / "empty-export.db")
    destination = tmp_path / f"empty{suffix}"
    count = export_rankings(
        db, "albums", destination, "2025-01-01", "2025-01-02", "daily", 50, "missing"
    )
    assert count == 0
    if suffix == ".xlsx":
        workbook = load_workbook(destination, read_only=True, data_only=True)
        assert workbook["Rankings"].max_row == 1
        definitions = list(workbook["Definitions"].values)
        assert ("Rows exported", "0") in definitions
    else:
        with destination.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))
        ranking_header_index = rows.index(
            next(row for row in rows if row and row[0] == "Rank")
        )
        assert len(rows) == ranking_header_index + 1
        assert ["Rows exported", "0"] in rows


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_export_escapes_spreadsheet_formula_text(tmp_path: Path, suffix: str) -> None:
    db = create_export_database(tmp_path / "formula-export.db")
    malicious_title = "=HYPERLINK(\"https://example.invalid\",\"open\")"
    db.register_chart_record(
        batch_id="formula-fixture",
        source="Authorized export fixture",
        source_license="Fixture license note",
        chart_type="daily",
        country="US",
        chart_date="2024-01-03",
        weekly_start=None,
        weekly_end=None,
        position=1,
        track_id="FORMULA_TRACK",
        track_name=malicious_title,
        artist_names="Formula Test",
        streams=10,
        original_url=None,
    )
    destination = tmp_path / f"formula{suffix}"
    assert export_rankings(
        db, "songs", destination, "2024-01-03", "2024-01-03", "daily", 50
    ) == 1
    if suffix == ".csv":
        with destination.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))
        header_index = next(index for index, row in enumerate(rows) if row and row[0] == "Rank")
        assert rows[header_index + 1][1].startswith("'=")
    else:
        workbook = load_workbook(destination, data_only=False)
        value = workbook["Rankings"].cell(2, 2)
        assert value.value.startswith("'=")
        assert value.data_type == "s"


def test_weekly_export_does_not_include_daily_streams(tmp_path: Path) -> None:
    db = create_export_database(tmp_path / "weekly-export.db")
    db.register_chart_record(
        batch_id="weekly-export-fixture",
        source="Authorized export fixture",
        source_license="Fixture license note",
        chart_type="weekly",
        country="US",
        chart_date="2024-01-01",
        weekly_start="2024-01-01",
        weekly_end="2024-01-07",
        position=1,
        track_id=TRACK_A,
        track_name="Song Alpha",
        artist_names="Credited Alpha",
        streams=900,
        original_url=None,
    )
    daily_path = tmp_path / "daily.csv"
    weekly_path = tmp_path / "weekly.csv"
    assert export_rankings(
        db, "songs", daily_path, "2024-01-01", "2024-01-01", "daily", 50, "Song Alpha"
    ) == 1
    assert export_rankings(
        db, "songs", weekly_path, "2024-01-01", "2024-01-01", "weekly", 50, "Song Alpha"
    ) == 1

    def total_from_csv(path: Path) -> int:
        with path.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))
        header_index = next(index for index, row in enumerate(rows) if row and row[0] == "Rank")
        return int(rows[header_index + 1][3])

    assert total_from_csv(daily_path) == 100
    assert total_from_csv(weekly_path) == 900
