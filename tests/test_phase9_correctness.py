"""Expected values come from Python fixture arithmetic, independently of SQL."""
from collections import defaultdict
import csv
import errno

import pytest
from openpyxl import load_workbook

from spotify_historical_analytics.db import DatabaseManager
from spotify_historical_analytics.exports import export_rankings
from spotify_historical_analytics.import_sessions import ImportSessionStore


ROWS = [
    ("2024-01-01", "t1", "Song Alpha", "Alice & Bob", 100, "daily", "US", "Oracle"),
    ("2024-01-02", "t1", "Song Alpha", "Alice & Bob", 120, "daily", "US", "Oracle"),
    ("2024-01-01", "t2", "Song Beta", "Alice", 80, "daily", "US", "Oracle"),
    ("2024-01-01", "t3", "Song Gamma", "Unknown credit", 50, "daily", "US", "Oracle"),
    ("2024-01-03", "t3", "Song Gamma", "Unknown credit", 70, "daily", "US", "Oracle"),
    ("2024-01-07", "t1", "Song Alpha", "Alice & Bob", 900, "weekly", "US", "Oracle"),
    ("2024-01-01", "t1", "Song Alpha", "Alice & Bob", 9999, "daily", "CA", "Oracle"),
    ("2024-01-01", "t1", "Song Alpha", "Alice & Bob", 30, "daily", "US", "DEMO / SYNTHETIC"),
]
ARTISTS = {"t1": ["a1", "a2"], "t2": ["a1"]}
ALBUMS = {"t1": ["edition-1", "edition-2"], "t2": ["edition-1"]}


def fixture_database(path):
    db = DatabaseManager(path)
    for day, track, name, credit, streams, cadence, country, source in ROWS:
        db.register_chart_record("oracle", source, "Synthetic test only", cadence, country,
                                 day, day if cadence == "weekly" else None,
                                 day if cadence == "weekly" else None, 1, track, name, credit, streams, None)
    # A repeated insertion must not add another observation.
    day, track, name, credit, streams, cadence, country, source = ROWS[0]
    db.register_chart_record("oracle", source, "Synthetic test only", cadence, country,
                             day, None, None, 1, track, name, credit, streams, None)
    db.connection.executemany("INSERT INTO artists(artist_id,artist_name) VALUES (?,?)", [("a1", "Alice"), ("a2", "Bob")])
    for track, artists in ARTISTS.items():
        for artist in artists:
            for rule in ("full-credit", "legacy-full-credit"):
                db.connection.execute("INSERT INTO track_artists(track_id,artist_id,attribution_rule,mapping_status) VALUES (?,?,?,'id_match')", (track, artist, rule))
    for album in ("edition-1", "edition-2"):
        db.connection.execute("INSERT INTO albums(album_id,album_name) VALUES (?, 'Same Album Name')", (album,))
    for track, albums in ALBUMS.items():
        for album in albums:
            db.connection.execute("INSERT INTO track_albums(track_id,album_id,mapping_status) VALUES (?,?,'id_match')", (track, album))
    db.connection.commit()
    return db


def oracle(start, end, cadence, include_demo):
    rows = [row for row in ROWS if row[6] == "US" and row[5] == cadence
            and start <= row[0] <= end and (include_demo or row[7] != "DEMO / SYNTHETIC")]
    songs, artists, albums, trend = (defaultdict(int) for _ in range(4))
    for day, track, name, credit, streams, *_ in rows:
        songs[track] += streams
        for artist in ARTISTS.get(track, ["credit:" + credit]):
            artists[artist] += streams
        for album in ALBUMS.get(track, []):
            albums[album] += streams
        trend[day] += streams
    return rows, songs, artists, albums, trend


@pytest.mark.parametrize("start,end,cadence", [
    ("2024-01-01", "2024-01-03", "daily"), ("2024-01-01", "2024-01-01", "daily"),
    ("2024-01-02", "2024-01-03", "daily"), ("2024-01-07", "2024-01-07", "weekly"),
    ("2023-01-01", "2023-01-02", "daily"),
])
@pytest.mark.parametrize("include_demo", [False, True])
def test_rankings_trends_collaborations_editions_and_boundaries(tmp_path, start, end, cadence, include_demo):
    db = fixture_database(tmp_path / "oracle.db")
    rows, songs, artists, albums, trend = oracle(start, end, cadence, include_demo)
    filters = dict(start_date=start, end_date=end, chart_type=cadence, include_demo=include_demo)
    summary = db.get_summary(**filters)
    assert summary["total_streams"] == sum(row[4] for row in rows)
    assert summary["distinct_tracks"] == len(songs)
    assert summary["distinct_artists"] == len({row[3] for row in rows})
    assert summary["chart_dates_with_data"] == len(trend)
    ranked = db.get_top_tracks(limit=200, **filters)
    assert {row["track_key"]: row["total_streams"] for row in ranked} == songs
    assert [r["total_streams"] for r in ranked] == sorted(songs.values(), reverse=True)
    actual = db.get_artist_summary(limit=200, **filters)
    assert {r["artist_id"] or r["artist_key"]: r["charting_song_streams"] for r in actual} == artists
    assert {r["album_id"]: r["charting_track_streams"] for r in db.get_album_rankings(limit=200, **filters)} == albums
    assert {r["chart_date"]: r["streams"] for r in db.get_streams_over_time(**filters)} == trend
    assert all("Alpha" in r["track_name"] for r in db.get_top_tracks(search="Alpha", **filters))
    assert all(r["artist_id"] == "a2" for r in db.get_artist_summary(search="Bob", **filters))
    db.close()


@pytest.mark.parametrize("category,column,index", [("songs", "Total chart streams", 3), ("artists", "Charting-song streams", 3), ("albums", "Album charting-track streams", 4)])
@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
@pytest.mark.parametrize("include_demo", [False, True])
def test_reopened_exports_match_independent_oracle(tmp_path, category, column, index, suffix, include_demo):
    db = fixture_database(tmp_path / "oracle.db")
    _, songs, artists, albums, _ = oracle("2024-01-01", "2024-01-03", "daily", include_demo)
    expected = sorted({"songs": songs, "artists": artists, "albums": albums}[category].values(), reverse=True)
    destination = tmp_path / (category + suffix)
    assert export_rankings(db, category, destination, "2024-01-01", "2024-01-03", "daily", 200, include_demo=include_demo) == len(expected)
    if suffix == ".csv":
        with destination.open(encoding="utf-8-sig", newline="") as handle:
            content = list(csv.reader(handle))
        header = next(i for i, row in enumerate(content) if row and row[0] == "Rank")
        assert content[header][index] == column
        assert [int(row[index]) for row in content[header + 1:] if row] == expected
        assert any("DEMO / SYNTHETIC DATA included" in " ".join(row) for row in content) == include_demo
    else:
        book = load_workbook(destination, data_only=True)
        assert [row[index] for row in list(book["Rankings"].values)[1:]] == expected
        metadata = dict(list(book["Definitions"].values)[1:])
        assert ("DEMO / SYNTHETIC" in metadata["Synthetic data"]) == include_demo
        assert book["Rankings"].cell(2, index + 1).number_format == "#,##0"
        book.close()
    db.close()


def test_failed_export_does_not_replace_existing_report(tmp_path, monkeypatch):
    db = fixture_database(tmp_path / "oracle.db")
    destination = tmp_path / "keep.csv"
    destination.write_text("Existing report")
    def fail(*args):
        raise OSError(errno.ENOSPC, "Simulated disk full")
    monkeypatch.setattr("spotify_historical_analytics.exports._write_csv", fail)
    with pytest.raises(OSError):
        export_rankings(db, "songs", destination, None, None, "daily", 50)
    assert destination.read_text() == "Existing report"
    assert not list(tmp_path.glob(".keep.*"))
    db.close()


@pytest.mark.parametrize("value,valid", [("9007199254740993", True), ("1,234", True), ("0", True), ("1,2", False), ("NaN", False), ("1.2", False), (str(2**63), False)])
def test_stream_integer_validation_is_exact(tmp_path, value, valid):
    source = tmp_path / "streams.csv"
    with source.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Date", "Track", "Artist", "Streams", "Region"])
        writer.writerow(["2024-01-01", "Synthetic", "Artist", value, "US"])
    store = ImportSessionStore(tmp_path / "stage.db")
    session = store.create_session(source, "Synthetic", "Test only", "daily")
    summary = store.prepare(session)
    assert summary.accepted_rows == int(valid)
    if valid:
        assert store.connection.execute("SELECT streams FROM staged_chart_rows").fetchone()[0] == int(value.replace(",", ""))
    store.close()
